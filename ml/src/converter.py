import base64
import io
import re
from typing import Any, Dict, List, Optional, Set, Tuple
import fitz  # PyMuPDF
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

# Word highlight color names -> CSS hex.
HIGHLIGHT_COLORS: Dict[str, str] = {
    "black": "#000000", "blue": "#0000FF", "cyan": "#00FFFF", "darkBlue": "#00008B",
    "darkCyan": "#008B8B", "darkGray": "#A9A9A9", "darkGreen": "#006400",
    "darkMagenta": "#8B008B", "darkRed": "#8B0000", "darkYellow": "#808000",
    "green": "#008000", "lightGray": "#D3D3D3", "magenta": "#FF00FF",
    "red": "#FF0000", "white": "#FFFFFF", "yellow": "#FFFF00",
}

EMU_PER_INCH = 914400
TWIPS_PER_PX = 96.0 / 1440.0

DEFAULT_PAGE_SETUP: Dict[str, Any] = {
    "page_size": "letter",
    "orientation": "portrait",
    "margins": {"top": 1, "bottom": 1, "left": 1.5, "right": 1},
}


def _color_to_hex(color_obj: Any) -> Optional[str]:
    """Extract RGB hex color string from a python-docx ColorFormat object."""
    if not color_obj:
        return None
    try:
        if color_obj.rgb:
            return f"#{color_obj.rgb}"
    except Exception:
        pass
    return None


def _get_paragraph_alignment(p: Paragraph) -> str:
    """Return text alignment CSS value for a paragraph."""
    align_map = {
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.LEFT: "left",
    }
    return align_map.get(p.alignment, "left")


def _highlight_css(run: Run) -> Optional[str]:
    """Return the CSS background-color for a run's highlight, if any."""
    try:
        rPr = run._element.rPr
        if rPr is None:
            return None
        hl_el = rPr.find(qn("w:highlight"))
        if hl_el is None:
            return None
        val = hl_el.get(qn("w:val"))
        if val and val != "none" and val in HIGHLIGHT_COLORS:
            return HIGHLIGHT_COLORS[val]
    except Exception:
        pass
    return None


def _run_to_html(run: Run) -> str:
    """Convert a single text run into styled HTML with font size, color, family, and styles."""
    text = run.text
    if not text:
        return ""

    escaped = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )

    styles: List[str] = []

    if run.font.size:
        pt_val = round(run.font.size.pt, 1)
        styles.append(f"font-size:{pt_val}pt")

    hex_color = _color_to_hex(run.font.color)
    if hex_color:
        styles.append(f"color:{hex_color}")

    if run.font.name:
        styles.append(f"font-family:'{run.font.name}', serif")

    res = escaped
    if styles:
        css_str = "; ".join(styles)
        res = f'<span style="{css_str}">{res}</span>'

    if run.bold:
        res = f"<strong>{res}</strong>"
    if run.italic:
        res = f"<em>{res}</em>"
    if run.underline:
        res = f"<u>{res}</u>"
    if run.font.strike:
        res = f"<s>{res}</s>"
    if run.font.subscript:
        res = f"<sub>{res}</sub>"
    if run.font.superscript:
        res = f"<sup>{res}</sup>"

    hl = _highlight_css(run)
    if hl:
        res = f'<mark style="background-color:{hl}">{res}</mark>'

    return res


def _image_html(blip: Any, doc_parts: Dict[str, bytes]) -> str:
    """Render an embedded drawing (blip) as a base64 data-URI <img>, preserving size."""
    embed_id = blip.get(qn("r:embed"))
    if not embed_id or embed_id not in doc_parts:
        return ""

    img_bytes = doc_parts[embed_id]
    b64_src = base64.b64encode(img_bytes).decode("utf-8")

    mime = "image/png"
    if img_bytes.startswith(b"\xff\xd8"):
        mime = "image/jpeg"
    elif img_bytes.startswith(b"GIF"):
        mime = "image/gif"

    width = None
    try:
        pic = blip.getparent().getparent()
        xfrm = pic.find(".//" + qn("a:xfrm")) if pic is not None else None
        ext = xfrm.find(qn("a:ext")) if xfrm is not None else None
        if ext is not None:
            cx = ext.get(qn("cx"))
            if cx:
                width = round(int(cx) / 9525)
    except Exception:
        pass

    if width:
        style = "max-width:100%; height:auto; display:block; border-radius:4px;"
        return f'<img src="data:{mime};base64,{b64_src}" width="{width}" style="{style}" />'
    else:
        style = "max-width:100%; height:auto; display:block; border-radius:4px;"

    return f'<img src="data:{mime};base64,{b64_src}" style="{style}" />'


def _run_or_image_html(run: Run, doc_parts: Dict[str, bytes]) -> str:
    """Render a run as either its inline image (if it contains a drawing) or styled text."""
    try:
        blips = run._element.findall(".//" + qn("a:blip"))
        if blips:
            return "".join(_image_html(b, doc_parts) for b in blips)
    except Exception:
        pass
    return _run_to_html(run)


def _is_dark_hex(hex_str: str) -> bool:
    """Check if a hex color is dark (luminance < 128) to ensure high text contrast."""
    if not hex_str:
        return False
    try:
        h = hex_str.lstrip("#")
        if len(h) == 6:
            r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
            lum = 0.299 * r + 0.587 * g + 0.114 * b
            return lum < 128
        if len(h) == 3:
            r, g, b = int(h[0] * 2, 16), int(h[1] * 2, 16), int(h[2] * 2, 16)
            lum = 0.299 * r + 0.587 * g + 0.114 * b
            return lum < 128
    except Exception:
        pass
    return False


def _paragraph_inner_html(p: Paragraph, doc_parts: Dict[str, bytes]) -> str:
    """Build the inner HTML (runs + hyperlinks + page breaks + images) of a paragraph."""
    parts: List[str] = []
    for child in p._element:
        tag = child.tag
        if tag == qn("w:r"):
            run = Run(child, p)
            brs = child.findall(qn("w:br"))
            page_breaks = ""
            for br in brs:
                if br.get(qn("w:type")) == "page":
                    page_breaks += '<hr class="page-break" data-page-break="true" />'
                else:
                    parts.append("<br/>")
            parts.append(page_breaks + _run_or_image_html(run, doc_parts))
        elif tag == qn("w:hyperlink"):
            parts.append(_hyperlink_html(child, p, doc_parts))
        elif tag == qn("w:br"):
            if child.get(qn("w:type")) == "page":
                parts.append('<hr class="page-break" data-page-break="true" />')
            else:
                parts.append("<br/>")
    return "".join(parts)


def _hyperlink_html(hyperlink_el: Any, p: Paragraph, doc_parts: Dict[str, bytes]) -> str:
    """Render a <w:hyperlink> element as an <a> tag with its styled run text."""
    try:
        r_id = hyperlink_el.get(qn("r:id"))
        url = "#"
        if r_id:
            try:
                url = p.part.rels[r_id].target_ref
            except Exception:
                url = "#"
        runs_html = "".join(
            _run_or_image_html(Run(r, p), doc_parts)
            for r in hyperlink_el.findall(qn("w:r"))
        )
        if not runs_html:
            return ""
        target = ' target="_blank" rel="noopener noreferrer"' if url.startswith("http") else ""
        return f'<a href="{url}"{target}>{runs_html}</a>'
    except Exception:
        return ""


def _paragraph_indent_css(p: Paragraph) -> str:
    """Return CSS for paragraph left indent / first-line indent (twips -> px)."""
    pPr = p._element.pPr
    if pPr is None:
        return ""
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        return ""
    styles: List[str] = []
    left = ind.get(qn("w:left")) or ind.get(qn("w:start"))
    if left and left != "0":
        styles.append(f"margin-left:{round(int(left) * TWIPS_PER_PX, 1)}px")
    first_line = ind.get(qn("w:firstLine"))
    if first_line and first_line != "0":
        styles.append(f"text-indent:{round(int(first_line) * TWIPS_PER_PX, 1)}px")
    hanging = ind.get(qn("w:hanging"))
    if hanging and hanging != "0":
        styles.append(f"text-indent:{round(-int(hanging) * TWIPS_PER_PX, 1)}px")
    return "; ".join(styles)


def _paragraph_spacing_css(p: Paragraph) -> str:
    """Return CSS for paragraph spacing from w:spacing (twips -> px, line -> unitless)."""
    pPr = p._element.pPr
    if pPr is None:
        return "margin-bottom:0.25em; line-height:1.5"
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        return "margin-bottom:0.25em; line-height:1.5"
    styles: List[str] = []
    before = sp.get(qn("w:before"))
    if before and before != "0":
        styles.append(f"margin-top:{round(int(before) * TWIPS_PER_PX, 1)}px")
    after = sp.get(qn("w:after"))
    if after and after != "0":
        styles.append(f"margin-bottom:{round(int(after) * TWIPS_PER_PX, 1)}px")
    else:
        styles.append("margin-bottom:0.25em")

    line = sp.get(qn("w:line"))
    if line and line != "0":
        rule = sp.get(qn("w:lineRule"), "auto")
        if rule == "exact" or rule == "atLeast":
            styles.append(f"line-height:{round(int(line) / 240.0, 2)}pt")
        else:
            styles.append(f"line-height:{round(int(line) / 240.0, 2)}")
    else:
        styles.append("line-height:1.5")

    return "; ".join(styles)


def _build_numbering_map(doc: docx.Document) -> Tuple[Dict[str, str], Dict[str, Dict[str, str]]]:
    """Return (numId->abstractNumId, abstractNumId->{ilvl: numFmt}) from numbering.xml."""
    try:
        numbering_part = doc.part.numbering_part
    except Exception:
        return {}, {}
    if numbering_part is None:
        return {}, {}

    root = numbering_part.element
    num_to_abstract: Dict[str, str] = {}
    for num in root.findall(qn("w:num")):
        num_id = num.get(qn("w:numId"))
        abstract = num.find(qn("w:abstractNumId"))
        if num_id and abstract is not None:
            num_to_abstract[num_id] = abstract.get(qn("w:val")) or ""

    abstract_levels: Dict[str, Dict[str, str]] = {}
    for abstract in root.findall(qn("w:abstractNum")):
        abstract_id = abstract.get(qn("w:abstractNumId"))
        if abstract_id is None:
            continue
        levels: Dict[str, str] = {}
        for lvl in abstract.findall(qn("w:lvl")):
            ilvl = lvl.get(qn("w:ilvl"), "0")
            fmt_el = lvl.find(qn("w:numFmt"))
            fmt = fmt_el.get(qn("w:val")) if fmt_el is not None else "bullet"
            levels[ilvl] = fmt
        abstract_levels[abstract_id] = levels

    return num_to_abstract, abstract_levels


def _paragraph_list_info(
    p: Paragraph,
    num_to_abstract: Dict[str, str],
    abstract_levels: Dict[str, Dict[str, str]],
) -> Optional[Dict[str, Any]]:
    """Return list kind + level for a numbered/bulleted paragraph, or None."""
    pPr = p._element.pPr
    if pPr is None:
        return None
    num_pr = pPr.find(qn("w:numPr"))
    if num_pr is None:
        return None
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        return None
    num_id = num_id_el.get(qn("w:val"))
    if not num_id or num_id == "0":
        return None

    ilvl_el = num_pr.find(qn("w:ilvl"))
    ilvl = int(ilvl_el.get(qn("w:val")) or "0") if ilvl_el is not None else 0

    fmt = "bullet"
    abstract_id = num_to_abstract.get(num_id)
    if abstract_id:
        levels = abstract_levels.get(abstract_id, {})
        fmt = levels.get(str(ilvl), levels.get("0", "bullet"))

    kind = "ul" if fmt == "bullet" else "ol"
    return {"kind": kind, "level": ilvl}


def _paragraph_to_html(
    p: Paragraph,
    doc_parts: Dict[str, bytes],
    list_info: Optional[Dict[str, Any]] = None,
) -> str:
    """Convert a paragraph to HTML; list paragraphs return only their inner content."""
    style_name = p.style.name.lower() if p.style else ""
    inner = _paragraph_inner_html(p, doc_parts)

    if list_info is not None:
        return inner

    if not inner.strip():
        # Preserve empty paragraphs so Word's blank-line spacing survives the
        # round trip instead of collapsing every paragraph together.
        return "<p><br/></p>"

    # A paragraph that is only a page break -> bare <hr>, not wrapped in <p>.
    if re.fullmatch(r"(<hr[^>]*/>\s*)+", inner):
        return inner

    # A paragraph that is only an image -> bare <img>, not wrapped in <p>.
    # This prevents ProseMirror from splitting the <p> and creating an
    # empty paragraph above the block-level image node.
    if re.fullmatch(r'<img\s[^>]+/>', inner.strip()):
        return inner.strip()

    alignment = _get_paragraph_alignment(p)
    align_style = f"text-align:{alignment};" if alignment != "left" else ""
    spacing_style = _paragraph_spacing_css(p)
    indent_style = _paragraph_indent_css(p)

    if "title" in style_name and "sub" not in style_name:
        css = " ".join(
            filter(None, [align_style, "font-size:20pt; font-weight:bold; margin:0.75em 0 0.35em;", indent_style])
        )
        return f'<h1 class="title" style="{css}">{inner}</h1>'
    if "subtitle" in style_name:
        css = " ".join(
            filter(None, [align_style, "font-size:14pt; font-weight:600; font-style:italic; margin:0.35em 0 0.5em;", indent_style])
        )
        return f'<h2 class="subtitle" style="{css}">{inner}</h2>'
    if "heading 1" in style_name:
        return f'<h1 style="{align_style} font-size:17pt; font-weight:bold; margin:0.85em 0 0.35em; {indent_style}">{inner}</h1>'
    if "heading 2" in style_name:
        return f'<h2 style="{align_style} font-size:14pt; font-weight:bold; margin:0.75em 0 0.3em; {indent_style}">{inner}</h2>'
    if "heading 3" in style_name:
        return f'<h3 style="{align_style} font-size:12.5pt; font-weight:bold; margin:0.65em 0 0.25em; {indent_style}">{inner}</h3>'
    if "heading 4" in style_name:
        return f'<h4 style="{align_style} font-size:11.5pt; font-weight:bold; margin:0.5em 0 0.25em; {indent_style}">{inner}</h4>'
    if "heading 5" in style_name:
        return f'<h5 style="{align_style} font-size:10.5pt; font-weight:bold; margin:0.5em 0 0.25em; {indent_style}">{inner}</h5>'
    if "heading 6" in style_name:
        return f'<h6 style="{align_style} font-size:9.5pt; font-weight:bold; margin:0.5em 0 0.25em; {indent_style}">{inner}</h6>'
    if "heading" in style_name:
        return f'<h4 style="{align_style} font-size:11.5pt; font-weight:bold; margin:0.5em 0 0.25em; {indent_style}">{inner}</h4>'

    css_parts = filter(None, [align_style, spacing_style, indent_style])
    css_combined = "; ".join(part.rstrip("; ") for part in css_parts).strip()
    style_attr = f' style="{css_combined};"' if css_combined else ""
    return f"<p{style_attr}>{inner}</p>"


def _cell_paragraph_to_html(p: Paragraph, doc_parts: Dict[str, bytes], force_white: bool = False) -> str:
    """Format paragraphs inside table cells compactly without ballooning cell heights."""
    inner = _paragraph_inner_html(p, doc_parts)
    if not inner.strip():
        return ""
    alignment = _get_paragraph_alignment(p)
    align_style = f"text-align:{alignment};" if alignment != "left" else ""
    color_style = "color:#ffffff;" if force_white else ""
    css = " ".join(filter(None, [align_style, color_style, "margin:0; line-height:1.35;"]))
    return f'<p style="{css}">{inner}</p>'


def _table_to_html(table: Table, doc_parts: Dict[str, bytes]) -> str:
    """Convert a Word table into styled HTML preserving column widths, borders, and cell padding."""
    tbl = table._element
    col_widths: List[int] = []
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc in grid.findall(qn("w:gridCol")):
            w = gc.get(qn("w:w"))
            if w:
                col_widths.append(int(w))
    total_width = sum(col_widths) or 1

    rows_html: List[str] = []
    for i, row in enumerate(table.rows):
        cells_html: List[str] = []
        tag = "th" if i == 0 else "td"

        for ci, cell in enumerate(row.cells):
            bg_color = ""
            text_color = ""
            is_dark = False
            try:
                tc_pr = cell._element.get_or_add_tcPr()
                shd = tc_pr.find(qn("w:shd"))
                if shd is not None and shd.get(qn("w:fill")):
                    fill = shd.get(qn("w:fill"))
                    if fill and fill != "auto" and fill.lower() not in ("ffffff", "none"):
                        bg_color = f"background-color:#{fill};"
                        if _is_dark_hex(fill):
                            is_dark = True
                            text_color = "color:#ffffff; font-weight:600;"
            except Exception:
                pass

            if not bg_color and i == 0:
                bg_color = "background-color:var(--secondary);"

            cell_content = "".join(_cell_paragraph_to_html(p, doc_parts, force_white=is_dark) for p in cell.paragraphs)
            if not cell_content.strip():
                cell_content = "&nbsp;"

            style_parts = ["border:1px solid var(--border)", "padding:8px 12px", "vertical-align:top"]
            if bg_color:
                style_parts.append(bg_color.rstrip(";"))
            if text_color:
                style_parts.append(text_color.rstrip(";"))
            if ci < len(col_widths):
                pct = col_widths[ci] / total_width * 100
                style_parts.append(f"width:{pct:.2f}%")

            cell_style = "; ".join(style_parts) + ";"
            cells_html.append(f'<{tag} style="{cell_style}">{cell_content}</{tag}>')

        tr_style = "background-color:var(--secondary);" if i == 0 else ""
        tr_attr = f' style="{tr_style}"' if tr_style else ""
        rows_html.append(f"<tr{tr_attr}>{''.join(cells_html)}</tr>")

    table_css = "border-collapse:collapse; width:100%; margin:1.25em 0; border:1px solid var(--border); table-layout:fixed;"
    return f'<table style="{table_css}">{"".join(rows_html)}</table>'


def _body_to_html(doc: docx.Document, doc_parts: Dict[str, bytes]) -> str:
    """Build the document body HTML, including properly nested lists."""
    num_to_abstract, abstract_levels = _build_numbering_map(doc)
    out: List[str] = []
    stack: List[Dict[str, str]] = []
    open_li_depth: Optional[int] = None

    def close_list() -> None:
        nonlocal open_li_depth
        if open_li_depth == len(stack):
            out.append("</li>")
            open_li_depth = None
        out.append(f"</{stack.pop()['kind']}>")
        # The <li> that contained this (nested) list is now complete.
        if len(stack) > 0:
            out.append("</li>")
            open_li_depth = None

    def close_all_lists() -> None:
        while stack:
            close_list()

    for element in doc.element.body:
        if element.tag.endswith("p"):
            p = Paragraph(element, doc)
            info = _paragraph_list_info(p, num_to_abstract, abstract_levels)
            if info is None:
                close_all_lists()
                html = _paragraph_to_html(p, doc_parts)
                if html:
                    out.append(html)
                continue

            kind = info["kind"]
            target_depth = info["level"] + 1

            # Close lists deeper than the target depth.
            while len(stack) > target_depth:
                close_list()
            # Same depth but different kind -> close and reopen as the new kind.
            if stack and len(stack) == target_depth and stack[-1]["kind"] != kind:
                close_list()

            if len(stack) == target_depth and stack:
                # Continuing an existing list at this depth -> new sibling item.
                if open_li_depth == target_depth:
                    out.append("</li>")
                out.append("<li>")
                open_li_depth = target_depth
            else:
                # Open nested list(s) (inside the current open <li>) or a fresh list.
                while len(stack) < target_depth:
                    stack.append({"kind": kind})
                    out.append(f"<{kind}>")
                out.append("<li>")
                open_li_depth = target_depth

            out.append(_paragraph_to_html(p, doc_parts, list_info=info))

        elif element.tag.endswith("tbl"):
            close_all_lists()
            out.append(_table_to_html(Table(element, doc), doc_parts))

    close_all_lists()
    return "\n".join(out)


def _extract_header_footer(doc: docx.Document) -> Tuple[str, str]:
    """Extract plain-text header and footer from the first section."""
    try:
        section = doc.sections[0] if doc.sections else None
    except Exception:
        return "", ""
    if section is None:
        return "", ""

    header = ""
    footer = ""
    try:
        if not section.header.is_linked_to_previous:
            header = " ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
    except Exception:
        pass
    try:
        if not section.footer.is_linked_to_previous:
            footer = " ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
    except Exception:
        pass
    return header, footer


def _page_size_key(short_in: float, long_in: float) -> str:
    """Map a section's short/long side (inches) to the closest editor page preset."""
    candidates = [
        ("letter", 8.5, 11.0),
        ("legal", 8.5, 14.0),
        ("a4", 8.27, 11.69),
        ("long", 8.5, 13.0),
        ("short", 8.5, 11.0),
    ]
    best = "letter"
    best_err: Optional[float] = None
    for key, short, long_side in candidates:
        err = abs(short_in - short) / short + abs(long_in - long_side) / long_side
        if best_err is None or err < best_err:
            best, best_err = key, err
    return best


def _extract_page_setup(doc: docx.Document) -> Dict[str, Any]:
    """Extract page size, orientation, and margins (inches) from the first section."""
    try:
        section = doc.sections[0] if doc.sections else None
        if section is None:
            return dict(DEFAULT_PAGE_SETUP)
    except Exception:
        return dict(DEFAULT_PAGE_SETUP)

    width_in = section.page_width / EMU_PER_INCH
    height_in = section.page_height / EMU_PER_INCH

    orientation = "landscape" if width_in > height_in else "portrait"
    short_side, long_side = (height_in, width_in) if width_in > height_in else (width_in, height_in)

    try:
        margins = {
            "top": round(section.top_margin / EMU_PER_INCH, 2),
            "bottom": round(section.bottom_margin / EMU_PER_INCH, 2),
            "left": round(section.left_margin / EMU_PER_INCH, 2),
            "right": round(section.right_margin / EMU_PER_INCH, 2),
        }
    except Exception:
        margins = dict(DEFAULT_PAGE_SETUP["margins"])

    return {
        "page_size": _page_size_key(short_side, long_side),
        "orientation": orientation,
        "margins": margins,
    }


def convert_docx_bytes_to_html(file_bytes: bytes) -> Dict[str, Any]:
    """High-fidelity DOCX to HTML converter using python-docx XML traversal.

    Returns html, page_setup, and plain-text header/footer so the server can
    persist page size/margins and reuse the editor's header/footer features.
    """
    doc_file = io.BytesIO(file_bytes)
    doc = docx.Document(doc_file)

    doc_parts: Dict[str, bytes] = {}
    for rel_id, rel in doc.part.rels.items():
        if "image" in rel.target_ref:
            try:
                doc_parts[rel_id] = rel.target_part.blob
            except Exception:
                pass

    html = _body_to_html(doc, doc_parts)
    header, footer = _extract_header_footer(doc)
    page_setup = _extract_page_setup(doc)

    return {
        "html": html,
        "page_setup": page_setup,
        "header": header,
        "footer": footer,
    }


# ─── PDF conversion (clean flowing HTML) ────────────────────────────────────

# PyMuPDF span flag bits (span["flags"]).
_PDF_FLAG_ITALIC = 2**1
_PDF_FLAG_BOLD = 2**4

# Top/bottom bands (fraction of page height) where running headers/footers live.
_PDF_FURNITURE_TOP = 0.07
_PDF_FURNITURE_BOTTOM = 0.93


def _pdf_furniture_band(line_y: float, page_height: float) -> Optional[str]:
    """Classify a line's vertical position as 'top', 'bottom', or None (body)."""
    normalized = line_y / page_height if page_height else 0.0
    if normalized < _PDF_FURNITURE_TOP:
        return "top"
    if normalized > _PDF_FURNITURE_BOTTOM:
        return "bottom"
    return None


# Matches page-number-like footer lines: "7", "Page 3", "3 of 12", "3/12".
_PDF_PAGENO_RE = re.compile(r"^(page\s*)?\d+(\s*(of|/)\s*\d+)?$", re.IGNORECASE)


def _detect_pdf_furniture(doc: Any) -> Set[Tuple[str, str]]:
    """Find running headers/footers repeated across pages so they can be dropped.

    Covers both identical repeated lines (company headers) and page-number
    footers whose digits change on every page.
    """
    page_count = len(doc)
    counters: Dict[Tuple[str, str], int] = {}
    band_texts: Dict[str, List[List[str]]] = {"top": [], "bottom": []}
    band_has_line: Dict[str, List[bool]] = {"top": [], "bottom": []}

    for page in doc:
        page_height = page.rect.height
        top_texts: List[str] = []
        bottom_texts: List[str] = []
        blocks = page.get_text("dict").get("blocks", [])
        for block in blocks:
            if block.get("type", 0) != 0:
                continue
            for line in block.get("lines", []):
                band = _pdf_furniture_band(line["bbox"][1], page_height)
                if band is None:
                    continue
                text = "".join(span.get("text", "") for span in line.get("spans", [])).strip()
                if not text:
                    continue
                counters[(band, text)] = counters.get((band, text), 0) + 1
                if band == "top":
                    top_texts.append(text)
                else:
                    bottom_texts.append(text)
        band_texts["top"].append(top_texts)
        band_texts["bottom"].append(bottom_texts)
        band_has_line["top"].append(bool(top_texts))
        band_has_line["bottom"].append(bool(bottom_texts))

    furniture: Set[Tuple[str, str]] = set()
    for band in ("top", "bottom"):
        required = max(2, round(page_count * 0.5))
        if sum(band_has_line[band]) < required:
            continue

        # Identical repeated lines.
        for (b, text), count in counters.items():
            if b == band and count >= required:
                furniture.add((band, text))

        # Page-number footers whose digits differ per page.
        pages_with_pageno = sum(
            1 for texts in band_texts[band] if any(_PDF_PAGENO_RE.match(t) for t in texts)
        )
        if pages_with_pageno >= required:
            for texts in band_texts[band]:
                for text in texts:
                    if _PDF_PAGENO_RE.match(text):
                        furniture.add((band, text))
    return furniture


def _pdf_span_to_html(span: Dict[str, Any]) -> str:
    """Render one PDF text span as HTML, preserving size/color/bold/italic."""
    text = span.get("text", "")
    if not text:
        return ""
    escaped = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )

    styles: List[str] = []
    size = span.get("size", 0)
    if size:
        styles.append(f"font-size:{round(size, 1)}pt")
    color = span.get("color")
    if color and color != 0:
        hex_color = f"#{color:06x}"
        styles.append(f"color:{hex_color}")

    result = f'<span style="{"; ".join(styles)}">{escaped}</span>' if styles else escaped

    flags = span.get("flags", 0)
    if flags & _PDF_FLAG_ITALIC:
        result = f"<em>{result}</em>"
    if flags & _PDF_FLAG_BOLD:
        result = f"<strong>{result}</strong>"
    return result


def _pdf_line_is_furniture(line: Dict[str, Any], page_height: float, furniture: Set[Tuple[str, str]]) -> bool:
    """Return True when a line is a repeated running header/footer to skip."""
    band = _pdf_furniture_band(line["bbox"][1], page_height)
    if band is None:
        return False
    text = "".join(span.get("text", "") for span in line.get("spans", [])).strip()
    return (band, text) in furniture


def _pdf_block_alignment(block: Dict[str, Any], page_width: float) -> Optional[str]:
    """Infer centered text: a short block whose center sits on the page midline."""
    x0, x1 = block["bbox"][0], block["bbox"][2]
    if x1 - x0 >= page_width * 0.6:
        return None
    center_x = (x0 + x1) / 2
    if abs(center_x - page_width / 2) < 15:
        return "center"
    return None


def _pdf_heading_level(block: Dict[str, Any]) -> Optional[str]:
    """Classify a single short bold line as an h1/h2 heading, or None."""
    lines = block.get("lines", [])
    if len(lines) != 1:
        return None
    spans = lines[0].get("spans", [])
    if not spans:
        return None
    sizes = [span.get("size", 0) for span in spans]
    max_size = max(sizes)
    all_bold = all(bool(span.get("flags", 0) & _PDF_FLAG_BOLD) for span in spans)
    text_len = len("".join(span.get("text", "") for span in spans).strip())
    if not all_bold or text_len == 0 or text_len > 100:
        return None
    if max_size >= 14:
        return "h1"
    if max_size >= 12:
        return "h2"
    return None


def _pdf_block_to_html(
    block: Dict[str, Any],
    page_width: float,
    page_height: float,
    furniture: Set[Tuple[str, str]],
) -> str:
    """Render a PDF text block as one flowing paragraph with inline styling."""
    lines_html: List[str] = []
    for line in block.get("lines", []):
        if _pdf_line_is_furniture(line, page_height, furniture):
            continue
        spans_html: List[str] = []
        for span in line.get("spans", []):
            html = _pdf_span_to_html(span)
            if html:
                spans_html.append(html)
        if spans_html:
            lines_html.append("".join(spans_html))
    if not lines_html:
        return ""

    inner = " ".join(lines_html)
    if not re.sub(r"<[^>]*>", "", inner).strip():
        # Whitespace-only blocks (stray blank lines / spacer spans) create
        # phantom empty paragraphs that push content onto blank pages.
        return ""
    alignment = _pdf_block_alignment(block, page_width)
    align_style = f"text-align:{alignment};" if alignment else ""
    heading = _pdf_heading_level(block)

    if heading == "h1":
        return f'<h1 style="{align_style} font-size:18pt; font-weight:bold; margin:0.75em 0 0.25em;">{inner}</h1>'
    if heading == "h2":
        return f'<h2 style="{align_style} font-size:15pt; font-weight:bold; font-style:italic; margin:0.75em 0 0.25em;">{inner}</h2>'
    if align_style:
        return f'<p style="{align_style}">{inner}</p>'
    return f"<p>{inner}</p>"


def _pdf_table_to_html(data: List[List[Optional[str]]]) -> str:
    """Convert extracted PDF table cells into styled HTML."""
    if not data:
        return ""
    rows_html: List[str] = []
    for i, row in enumerate(data):
        cells_html: List[str] = []
        for cell in row:
            cell_text = (cell or "").strip().replace("\n", " ")
            escaped = (
                cell_text.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            tag = "th" if i == 0 else "td"
            content = escaped if escaped else "&nbsp;"
            cells_html.append(f'<{tag} style="border:1px solid var(--border); padding:6px 10px;">{content}</{tag}>')
        tr_style = "background-color:var(--secondary);" if i == 0 else ""
        tr_attr = f' style="{tr_style}"' if tr_style else ""
        rows_html.append(f"<tr{tr_attr}>{''.join(cells_html)}</tr>")
    table_css = "border-collapse:collapse; width:100%; margin:1em 0; border:1px solid var(--border); table-layout:fixed;"
    return f'<table style="{table_css}">{"".join(rows_html)}</table>'


MAX_IMAGE_DIM = 1400
JPEG_QUALITY = 82


def _compress_image_bytes(img_bytes: bytes) -> Tuple[bytes, str]:
    """Downscale an embedded image to MAX_IMAGE_DIM and re-encode it small.

    Returns (new_bytes, mime). Falls back to the original bytes + 'image/png'
    if decoding or re-encoding fails, so conversion never breaks on odd images.
    """
    try:
        pix = fitz.Pixmap(img_bytes)
        largest = max(pix.width, pix.height)
        if largest > MAX_IMAGE_DIM:
            factor = MAX_IMAGE_DIM / largest
            pix = fitz.Pixmap(
                pix, int(pix.width * factor), int(pix.height * factor)
            )
        if pix.alpha:
            out_bytes = pix.tobytes("png")
            return out_bytes, "image/png"
        if pix.colorspace is not None and pix.colorspace.n > 3:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        out_bytes = pix.tobytes("jpeg", jpg_quality=JPEG_QUALITY)
        return out_bytes, "image/jpeg"
    except Exception:
        return img_bytes, "image/png"


def _pdf_image_to_html(doc: Any, info: Dict[str, Any]) -> str:
    """Render a PDF image as a compressed inline base64 <img>, preserving its width."""
    try:
        xref = info.get("xref")
        if not xref:
            return ""
        extracted = doc.extract_image(xref)
        img_bytes, mime = _compress_image_bytes(extracted["image"])
        b64 = base64.b64encode(img_bytes).decode("utf-8")
        bbox = info.get("bbox")
        width = round(bbox[2] - bbox[0]) if bbox else None
        if width:
            style = "max-width:100%; height:auto; display:block; border-radius:4px;"
            return f'<img src="data:{mime};base64,{b64}" width="{width}" style="{style}" />'
        else:
            style = "max-width:100%; height:auto; display:block; border-radius:4px;"
        return f'<img src="data:{mime};base64,{b64}" style="{style}" />'
    except Exception:
        return ""


def _pdf_page_to_html(page: Any, doc: Any, furniture: Set[Tuple[str, str]]) -> str:
    """Convert one PDF page into clean flowing HTML (text + images + tables)."""
    page_width = page.rect.width
    page_height = page.rect.height

    tables: List[Tuple[float, str]] = []
    table_boxes: List[Tuple[float, float, float, float]] = []
    try:
        for table in page.find_tables():
            bbox = table.bbox
            html = _pdf_table_to_html(table.extract())
            if html:
                tables.append((bbox[1], html))
                table_boxes.append((bbox[0], bbox[1], bbox[2], bbox[3]))
    except Exception:
        pass

    def inside_table(bbox: Tuple[float, float, float, float]) -> bool:
        for tb in table_boxes:
            if bbox[0] >= tb[0] - 2 and bbox[1] >= tb[1] - 2 and bbox[2] <= tb[2] + 2 and bbox[3] <= tb[3] + 2:
                return True
        return False

    items: List[Tuple[float, str]] = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type", 0) != 0:
            continue
        if inside_table(tuple(block["bbox"])):
            continue
        html = _pdf_block_to_html(block, page_width, page_height, furniture)
        if html:
            items.append((block["bbox"][1], html))

    for info in page.get_image_info(xrefs=True):
        bbox = info.get("bbox")
        if not bbox:
            continue
        img_html = _pdf_image_to_html(doc, info)
        if img_html:
            items.append((bbox[1], img_html))

    items.extend(tables)
    items.sort(key=lambda item: item[0])
    return "".join(html for _, html in items)


def _pdf_page_setup(page: Any) -> Dict[str, Any]:
    """Derive page size and orientation from the first PDF page's dimensions."""
    width_in = page.rect.width / 72.0
    height_in = page.rect.height / 72.0
    orientation = "landscape" if width_in > height_in else "portrait"
    short_side, long_side = (height_in, width_in) if width_in > height_in else (width_in, height_in)
    return {
        "page_size": _page_size_key(short_side, long_side),
        "orientation": orientation,
        "margins": dict(DEFAULT_PAGE_SETUP["margins"]),
    }


def convert_pdf_bytes_to_html(file_bytes: bytes) -> Dict[str, Any]:
    """Convert a PDF to clean flowing HTML: paragraphs with bold/italic/size,
    plus embedded figures and tables, dropping repeated running headers/footers."""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    furniture = _detect_pdf_furniture(doc)
    html_parts: List[str] = []
    for page in doc:
        html_parts.append(_pdf_page_to_html(page, doc, furniture))
    page_setup = _pdf_page_setup(doc[0]) if len(doc) else dict(DEFAULT_PAGE_SETUP)
    doc.close()
    return {
        "html": "\n".join(html_parts),
        "page_setup": page_setup,
        "header": "",
        "footer": "",
    }
