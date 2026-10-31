"""Alter ~10% of the prose paragraphs of the clean V2 document and record which ones.

Writes tests/fixtures/import_baselines/documents/V2 INTELLIDOCS light sabotage.docx plus a .json list of altered texts.
"""
import json
import random
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DOCS = Path(__file__).resolve().parents[2] / "fixtures" / "import_baselines" / "documents"
doc = docx.Document(DOCS / "V2 INTELLIDOCS (2).docx")
paras = [p for p in doc.paragraphs if len(p.text.strip()) >= 80 and p.runs]
start = next((i for i, p in enumerate(doc.paragraphs) if p.text.strip().upper() == "INTRODUCTION"), 0)
body = [p for p in paras if any(p._p is q._p for q in doc.paragraphs[start:])]
random.seed(7)
picked = random.sample(body, max(6, len(body) // 10))
kinds = ["bold", "italic", "left", "size"]
log = []
for i, p in enumerate(picked):
    kind = kinds[i % 4]
    if kind == "bold":
        for r in p.runs: r.bold = True
    elif kind == "italic":
        for r in p.runs: r.italic = True
    elif kind == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        for r in p.runs: r.font.size = Pt(16)
    log.append({"text": p.text.strip(), "kind": kind})
doc.save(DOCS / "V2 INTELLIDOCS light sabotage.docx")
(DOCS / "V2 INTELLIDOCS light sabotage.json").write_text(json.dumps(log, indent=1), encoding="utf-8")
print(len(body), "body paragraphs;", len(log), "altered")
