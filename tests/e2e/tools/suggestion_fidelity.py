"""Run the real formatting scanner on imported documents and score it against the clean/sabotaged pair.

Usage: python tests/e2e/tools/suggestion_fidelity.py --web http://localhost:5183 [--preset uclm_capstone]
Ground truth = paragraphs whose alignment, bold, italic, size or spacing differ between the clean and the sabotaged DOCX.
"""
import argparse
import json
import sys
from pathlib import Path

import docx
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "tests" / "e2e"))
sys.path.insert(0, str(ROOT / "ml" / "src"))
import fidelity_run as fr  # noqa: E402

DOCS = ROOT / "tests" / "fixtures" / "import_baselines" / "documents"
CLEAN, BAD = "V2 INTELLIDOCS (2).docx", "V2 INTELLIDOCS sabotage.docx"


def props(p):
    """The formatting facts of one paragraph that the sabotage may have changed."""
    runs = [r for r in p.runs if r.text.strip()]
    return {
        "align": str(p.alignment),
        "bold": any(r.bold for r in runs),
        "italic": any(r.italic for r in runs),
        "size": next((r.font.size.pt for r in runs if r.font.size), None),
        "space": (p.paragraph_format.space_before, p.paragraph_format.space_after),
    }


def ground_truth():
    """text -> list of changed property names, for every paragraph that differs."""
    a = [p for p in docx.Document(DOCS / CLEAN).paragraphs if p.text.strip()]
    b = [p for p in docx.Document(DOCS / BAD).paragraphs if p.text.strip()]
    changed = {}
    for x, y in zip(a, b):
        px, py = props(x), props(y)
        diff = [k for k in px if px[k] != py[k]]
        if diff:
            changed[y.text.strip()] = diff
    return changed


def scan(browser, source, web, preset):
    """The scanner's suggestions for one imported file."""
    payload = fr.to_payload(fr.convert(DOCS / source))
    page = browser.new_context(viewport={"width": 1300, "height": 1250}).new_page()
    page.route("**fonts.g*.com**", lambda r: r.abort())
    page.add_init_script("window.__FIDELITY__ = " + json.dumps(payload))
    q = "scan=1" + (f"&preset={preset}" if preset else "")
    page.goto(f"{web}/dev/editor-harness?{q}", wait_until="domcontentloaded")
    page.wait_for_selector(".rm-pages-wrapper", state="attached", timeout=90000)
    page.wait_for_timeout(14000)
    return page.evaluate("window.__scanner || []")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--web", default="http://localhost:5183")
    ap.add_argument("--preset", default=None)
    args = ap.parse_args()
    light = {e["text"]: [e["kind"]] for e in json.loads((DOCS / "V2 INTELLIDOCS light sabotage.json").read_text(encoding="utf-8"))}
    cases = [(CLEAN, {}), ("V2 INTELLIDOCS light sabotage.docx", light), (BAD, ground_truth())]
    with sync_playwright() as pw:
        browser = pw.chromium.launch(channel="chrome")
        for name, truth in cases:
            sug = scan(browser, name, args.web, args.preset)
            hit = {t for t in truth if any(x["preview"].strip() and t.startswith(x["preview"].strip()) for x in sug)}
            wrong = [x for x in sug if not any(t.startswith(x["preview"].strip()) for t in truth)]
            print(f"{name}: {len(sug)} suggestions; altered {len(truth)}; found {len(hit)}; unwanted {len(wrong)}")
            for x in wrong[:6]:
                print(f"   unwanted {x['format']:<10} {x['confidence']:>3} {x['preview']!r}")
        browser.close()


if __name__ == "__main__":
    main()
