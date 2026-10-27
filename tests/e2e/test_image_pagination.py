"""Images and tables at page boundaries, in the paged editor (dev harness: no login, no database).

Regressions covered:
  * an image that does not fit the space left on a page used to collapse to 0x0 and vanish, because the
    resize library's flex node view cannot sit beside the paginator's full-width page-break float;
  * an image taller than a page body fits on no page, so the paginator added pages forever
    (the cap in styles/index.css keeps it to one page).

Needs the frontend dev server (WEB_URL, default http://localhost:5173) and installed Chrome.
"""
import io
import json
import os
import sys
import urllib.request
from pathlib import Path

import cv2
import docx
import numpy as np
import pytest
from docx.shared import Inches
from playwright.sync_api import sync_playwright

sys.path.insert(0, str(Path(__file__).resolve().parent))
import fidelity_run as fr  # noqa: E402

WEB = os.environ.get("WEB_URL", "http://localhost:5173").rstrip("/")
BODY_HEIGHT_PX = 864  # letter page, 1in margins


def _frontend_up() -> bool:
    try:
        urllib.request.urlopen(WEB, timeout=3)
        return True
    except OSError:
        return False


pytestmark = pytest.mark.skipif(not _frontend_up(), reason=f"blocked: frontend not running at {WEB}")

MEASURE = """() => {
  const pm = document.querySelector('.ProseMirror');
  const box = pm.getBoundingClientRect();
  const wrapper = document.querySelector('.rm-pages-wrapper');
  const pageH = parseFloat(getComputedStyle(wrapper).getPropertyValue('--rm-page-height'));
  const pitch = pageH + 40;
  const rect = (el) => { const r = el.getBoundingClientRect(); return {w: Math.round(r.width), h: Math.round(r.height),
    page: Math.floor((r.top - box.top + 1) / pitch) + 1, y: Math.round((r.top - box.top) % pitch)}; };
  return {pages: Math.round((box.height + 40) / pitch), imgs: Array.from(pm.querySelectorAll('img')).map(rect),
          cells: Array.from(pm.querySelectorAll('td')).map(rect)};
}"""


@pytest.fixture(scope="module")
def browser():
    with sync_playwright() as pw:
        instance = pw.chromium.launch(headless=True, channel="chrome", args=["--font-render-hinting=none"])
        yield instance
        instance.close()


def build_docx(tmp_path: Path, lines: int, picture: tuple[float, float] | None = None, table: bool = False) -> Path:
    """A letter DOCX: `lines` short paragraphs, then an optional picture (inches) and/or a 3x2 table."""
    document = docx.Document()
    for section in document.sections:
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)
    for index in range(lines):
        document.add_paragraph(f"Line {index}")
    if picture:
        width_in, height_in = picture
        pixels = np.full((max(1, int(height_in * 60)), int(width_in * 60), 3), 200, dtype=np.uint8)
        ok, encoded = cv2.imencode(".png", pixels)
        assert ok
        document.add_picture(io.BytesIO(encoded.tobytes()), width=Inches(width_in), height=Inches(height_in))
    if table:
        grid = document.add_table(rows=3, cols=2)
        grid.style = "Table Grid"
        for row in range(3):
            for col in range(2):
                grid.cell(row, col).text = f"r{row}c{col}"
    document.add_paragraph("after")
    path = tmp_path / "case.docx"
    document.save(path)
    return path


def measure(browser, path: Path, settle_ms: int = 4000) -> dict:
    """Import the DOCX with the real converter, lay it out in the harness and measure it."""
    payload = fr.to_payload(fr.convert(path))
    page = browser.new_context(viewport={"width": 1300, "height": 1250}).new_page()
    page.route("**fonts.googleapis.com**", lambda route: route.abort())
    page.route("**fonts.gstatic.com**", lambda route: route.abort())
    page.add_init_script("window.__FIDELITY__ = " + json.dumps(payload))
    page.goto(f"{WEB}/dev/editor-harness", wait_until="domcontentloaded")
    page.wait_for_selector(".rm-pages-wrapper", state="attached", timeout=30000)
    page.wait_for_timeout(settle_ms)
    first = page.evaluate(MEASURE)
    page.wait_for_timeout(2000)
    second = page.evaluate(MEASURE)
    page.context.close()
    assert first["pages"] == second["pages"], f"layout still changing: {first['pages']} -> {second['pages']} pages"
    return second


def test_image_that_fits_stays_on_its_page(browser, tmp_path: Path) -> None:
    result = measure(browser, build_docx(tmp_path, lines=10, picture=(3, 3)))
    assert result["imgs"] == [{"w": 288, "h": 288, "page": 1, "y": result["imgs"][0]["y"]}]
    assert result["pages"] == 1


@pytest.mark.parametrize("lines", [17, 20, 24])
def test_image_that_does_not_fit_moves_to_the_next_page(browser, tmp_path: Path, lines: int) -> None:
    """A 3in image after `lines` lines straddles the page end: it must reappear at the top of page 2."""
    result = measure(browser, build_docx(tmp_path, lines=lines, picture=(3, 3)))
    (image,) = result["imgs"]
    assert (image["w"], image["h"]) == (288, 288), f"image collapsed: {image}"
    assert image["page"] == 2
    assert 90 <= image["y"] <= 100  # top margin
    assert result["pages"] == 2


@pytest.mark.parametrize("height_in", [9.5, 12])
def test_image_taller_than_a_page_is_capped_and_pagination_settles(browser, tmp_path: Path, height_in: float) -> None:
    result = measure(browser, build_docx(tmp_path, lines=1, picture=(3, height_in)))
    (image,) = result["imgs"]
    assert 0 < image["h"] <= BODY_HEIGHT_PX, f"image not capped: {image}"
    assert result["pages"] <= 4  # before the cap this grew without bound


@pytest.mark.parametrize("lines", [18, 22, 26])
def test_table_near_a_page_bottom_keeps_its_cells(browser, tmp_path: Path, lines: int) -> None:
    result = measure(browser, build_docx(tmp_path, lines=lines, table=True))
    assert len(result["cells"]) == 6
    assert all(cell["w"] > 0 and cell["h"] > 0 for cell in result["cells"])


def open_in_harness(browser, path: Path):
    """Import the DOCX and return the harness page once it has laid out."""
    payload = fr.to_payload(fr.convert(path))
    page = browser.new_context(viewport={"width": 1500, "height": 1250}).new_page()
    page.route("**fonts.googleapis.com**", lambda route: route.abort())
    page.route("**fonts.gstatic.com**", lambda route: route.abort())
    page.add_init_script("window.__FIDELITY__ = " + json.dumps(payload))
    page.goto(f"{WEB}/dev/editor-harness", wait_until="domcontentloaded")
    page.wait_for_selector(".rm-pages-wrapper", state="attached", timeout=30000)
    page.wait_for_timeout(4000)
    return page


def test_dragging_an_image_past_the_margin_stores_the_width_it_draws(browser, tmp_path: Path) -> None:
    """The stored width is capped to the text column (624px here), so saved size == drawn size."""
    page = open_in_harness(browser, build_docx(tmp_path, lines=1, picture=(3, 1.5)))
    size = "() => { const i = document.querySelector('.ProseMirror img'); return {drawn: Math.round(i.getBoundingClientRect().width), stored: parseFloat(i.style.width)}; }"
    page.click(".ProseMirror img")
    page.wait_for_timeout(400)
    for drag in (400, 700):
        box = page.locator('div[style*="nwse-resize"]').last.bounding_box()
        x, y = box["x"] + box["width"] / 2, box["y"] + box["height"] / 2
        page.mouse.move(x, y)
        page.mouse.down()
        page.mouse.move(x + drag / 2, y, steps=5)
        page.mouse.move(x + drag, y, steps=5)
        page.mouse.up()
        page.wait_for_timeout(600)
        page.click(".ProseMirror img")
        page.wait_for_timeout(300)
    result = page.evaluate(size)
    page.context.close()
    assert abs(result["drawn"] - 624) <= 2, result  # the selected image has a 1px border each side
    assert result["stored"] <= 624, f"stored width {result['stored']} is wider than the text column"


def test_header_and_footer_pictures_fit_inside_the_margin(browser, tmp_path: Path) -> None:
    """A 0.8in header logo and 0.9in footer banner used to overlap the body; now they stay in the margin."""
    document = docx.Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.5)
    pixels = np.full((120, 300, 3), 200, dtype=np.uint8)
    ok, encoded = cv2.imencode(".png", pixels)
    assert ok
    section.header.paragraphs[0].add_run().add_picture(io.BytesIO(encoded.tobytes()), width=Inches(2), height=Inches(0.8))
    section.footer.paragraphs[0].add_run().add_picture(io.BytesIO(encoded.tobytes()), width=Inches(3), height=Inches(0.9))
    document.add_paragraph("Body text")
    path = tmp_path / "hf.docx"
    document.save(path)

    page = open_in_harness(browser, path)
    zones = page.evaluate("""() => {
      const box = document.querySelector('.ProseMirror').getBoundingClientRect();
      const first = (sel) => Array.from(document.querySelectorAll(sel + ' img')).map(i => i.getBoundingClientRect()).find(r => r.width > 0 && r.top - box.top < 1056);
      const head = first('.rm-page-header'), foot = first('.rm-page-footer');
      return {headBottom: Math.round(head.bottom - box.top), footTop: Math.round(foot.top - box.top)};
    }""")
    page.context.close()
    assert zones["headBottom"] <= 96 + 1, f"header picture reaches into the body: {zones}"
    assert zones["footTop"] >= 960 - 1, f"footer picture reaches into the body: {zones}"


def test_list_bullet_style_renders_as_list_items(browser, tmp_path: Path) -> None:
    document = docx.Document()
    for index in range(3):
        document.add_paragraph(f"bullet {index}", style="List Bullet")
    page = open_in_harness(browser, _save(document, tmp_path / "list.docx"))
    items = page.evaluate("() => Array.from(document.querySelectorAll('.ProseMirror li')).map(li => Math.round(li.getBoundingClientRect().height))")
    page.context.close()
    assert len(items) == 3 and all(height > 0 for height in items)


def _save(document, path: Path) -> Path:
    """Save a python-docx document and return its path."""
    document.save(path)
    return path
