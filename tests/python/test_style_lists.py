"""Tests that Word's style-based lists ("List Bullet", "List Number") import as real lists."""
import io
import os
import re
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "ml", "src"))

from converter import convert_docx_bytes_to_html  # noqa: E402


def to_html(document: docx.document.Document) -> str:
    """Convert an in-memory DOCX with the real converter."""
    buffer = io.BytesIO()
    document.save(buffer)
    return convert_docx_bytes_to_html(buffer.getvalue())["html"]


def item_texts(html: str) -> list[str]:
    """The text of every <li>, in order."""
    return [re.sub(r"<[^>]+>", "", item).strip() for item in re.findall(r"<li\b.*?</li>", html, flags=re.S)]


def test_list_bullet_style_becomes_a_bulleted_list() -> None:
    document = docx.Document()
    document.add_paragraph("Intro")
    for index in range(3):
        document.add_paragraph(f"bullet {index}", style="List Bullet")
    document.add_paragraph("Outro")
    html = to_html(document)
    assert html.count("<ul") == 1 and html.count("<ol") == 0
    assert item_texts(html) == ["bullet 0", "bullet 1", "bullet 2"]
    assert html.index("Intro") < html.index("<ul") < html.index("bullet 0") < html.index("</ul>") < html.index("Outro")


def test_list_number_style_becomes_a_numbered_list() -> None:
    document = docx.Document()
    for index in range(2):
        document.add_paragraph(f"step {index}", style="List Number")
    html = to_html(document)
    assert html.count("<ol") == 1 and html.count("<ul") == 0
    assert item_texts(html) == ["step 0", "step 1"]


def test_plain_paragraphs_stay_paragraphs() -> None:
    document = docx.Document()
    document.add_paragraph("one")
    document.add_paragraph("two")
    assert "<li" not in to_html(document)


def test_paragraph_that_switches_numbering_off_stays_plain() -> None:
    """A direct numId of 0 beats the numbering its style would give."""
    document = docx.Document()
    document.add_paragraph("kept", style="List Bullet")
    plain = document.add_paragraph("not a bullet", style="List Bullet")
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.append(num_id)
    plain._p.get_or_add_pPr().append(num_pr)
    assert item_texts(to_html(document)) == ["kept"]


def test_numbered_heading_style_is_not_turned_into_a_list_item() -> None:
    """Outline numbering on a heading style ("1.1 Scope") stays a heading."""
    document = docx.Document()
    heading_pPr = document.styles["Heading 1"].element.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.append(num_id)
    heading_pPr.append(num_pr)
    document.add_heading("Introduction", level=1)
    html = to_html(document)
    assert "<li" not in html
    assert re.search(r"<h1\b", html)
