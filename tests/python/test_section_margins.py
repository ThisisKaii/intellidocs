"""Tests that every Word section keeps its own left/right margins in the imported HTML."""
import io
import os
import sys

import docx
from docx.enum.section import WD_SECTION
from docx.shared import Inches

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "ml", "src"))

from converter import convert_docx_bytes_to_html  # noqa: E402


def build_document(left_margins: list[float]) -> bytes:
    """A DOCX with one section per entry, each holding one paragraph and its own left margin (inches)."""
    document = docx.Document()
    for index, left in enumerate(left_margins):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION.NEW_PAGE)
        section.left_margin = Inches(left)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        document.add_paragraph(f"Paragraph of section {index}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def paragraph_tag(html: str, text: str) -> str:
    """The opening <p ...> tag of the paragraph holding the given text."""
    end = html.index(text)
    return html[html.rindex("<p", 0, end):end]


def test_odd_section_is_shifted_by_its_margin_difference() -> None:
    """Sections that share the document-wide margin stay put; the odd one moves by the difference."""
    html = convert_docx_bytes_to_html(build_document([1.0, 2.0, 1.0]))["html"]
    assert "margin-left:96px" in paragraph_tag(html, "Paragraph of section 1")
    assert "margin-left:96px" not in paragraph_tag(html, "Paragraph of section 0")
    assert "margin-left:96px" not in paragraph_tag(html, "Paragraph of section 2")


def build_with_blocks(left_margins: list[float]) -> bytes:
    """Like build_document, plus a one-cell table and a numbered item in every section."""
    document = docx.Document()
    for index, left in enumerate(left_margins):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION.NEW_PAGE)
        section.left_margin = Inches(left)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        document.add_paragraph(f"Paragraph of section {index}")
        document.add_table(rows=1, cols=1).cell(0, 0).text = f"Cell of section {index}"
        document.add_paragraph(f"Item of section {index}", style="List Number")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def table_tag(html: str, text: str) -> str:
    """The opening <table ...> tag of the table holding the given text."""
    end = html.index(text)
    return html[html.rindex("<table", 0, end):html.index(">", html.rindex("<table", 0, end))]


def test_table_in_an_odd_section_carries_its_inset() -> None:
    """A table moves with its section: left, right and top offsets travel as data-inset."""
    html = convert_docx_bytes_to_html(build_with_blocks([1.0, 2.0, 1.0]))["html"]
    assert 'data-inset="96,0,0"' in table_tag(html, "Cell of section 1")
    assert "data-inset" not in table_tag(html, "Cell of section 0")
    assert "data-inset" not in table_tag(html, "Cell of section 2")


def test_list_marker_moves_with_its_section() -> None:
    """A numbered item's marker is placed at an absolute x, so it shifts by the same margin difference."""
    html = convert_docx_bytes_to_html(build_with_blocks([1.0, 2.0, 1.0]))["html"]

    def marker_x(text: str) -> float:
        item = html[html.rindex("<li", 0, html.index(text)):html.index(text)]
        return float(item.split('data-mx="')[1].split('"')[0])

    assert marker_x("Item of section 1") - marker_x("Item of section 0") == 96
    assert marker_x("Item of section 2") == marker_x("Item of section 0")


def test_later_section_is_shifted_not_only_the_first() -> None:
    """A narrower margin in the last section is applied too (negative shift)."""
    html = convert_docx_bytes_to_html(build_document([1.0, 1.0, 0.5]))["html"]
    assert "margin-left:-48px" in paragraph_tag(html, "Paragraph of section 2")
