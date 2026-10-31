"""A cell merged down several rows is laid out across them, not counted twice; imported lists carry no margins."""
import io
import os
import re
import sys
from typing import List, Optional

import docx
from lxml import etree

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "ml", "src"))

from converter import convert_docx_bytes_to_html  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def cell(text_paragraphs: List[str], merge: Optional[str] = None, width: int = 4000, before: int = 0) -> str:
    """One w:tc holding the given paragraphs (8pt text); `merge` is "restart" or "continue" for a vertical merge."""
    merge_xml = {None: "", "restart": '<w:vMerge w:val="restart"/>', "continue": "<w:vMerge/>"}[merge]
    paragraphs = "".join(
        f'<w:p><w:pPr><w:spacing w:before="{before if index else 0}" w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        f'<w:r><w:rPr><w:sz w:val="16"/></w:rPr><w:t>{text}</w:t></w:r></w:p>'
        for index, text in enumerate(text_paragraphs)
    ) or "<w:p/>"
    return f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{merge_xml}</w:tcPr>{paragraphs}</w:tc>'


def row(cells: List[str], height: Optional[int] = None) -> str:
    height_xml = f'<w:trPr><w:trHeight w:val="{height}" w:hRule="exact"/></w:trPr>' if height else ""
    return f"<w:tr>{height_xml}{''.join(cells)}</w:tr>"


def convert(rows: List[str]) -> str:
    document = docx.Document()
    table = etree.fromstring(
        f'<w:tbl xmlns:w="{W}"><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="4000"/></w:tblGrid>' + "".join(rows) + "</w:tbl>"
    )
    sect = document.element.body.find(f"{{{W}}}sectPr")
    sect.addprevious(table)
    buffer = io.BytesIO()
    document.save(buffer)
    return convert_docx_bytes_to_html(buffer.getvalue())["html"]


def cells_of_rows(html: str) -> List[List[str]]:
    """The visible text of every cell of every row."""
    rows = re.findall(r"<tr>(.*?)</tr>", html, flags=re.S)
    return [
        [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", c).replace("&nbsp;", "")).strip() for c in re.findall(r"<td[^>]*>(.*?)</td>", r, flags=re.S)]
        for r in rows
    ]


def test_merged_text_is_spread_over_the_rows_it_covers_by_where_it_starts() -> None:
    # Row 0 is 20px tall (300 twips); the second paragraph starts below it, so it belongs to row 1.
    html = convert([
        row([cell(["Label"]), cell(["First line", "Second line"], "restart", before=240)], height=300),
        row([cell(["Other"]), cell([], "continue")], height=600),
    ])
    assert cells_of_rows(html) == [["Label", "First line"], ["Other", "Second line"]]


def test_a_paragraph_that_fits_in_the_first_row_stays_there() -> None:
    html = convert([
        row([cell(["Label"]), cell(["One", "Two"], "restart")], height=2000),
        row([cell(["Other"]), cell([], "continue")], height=400),
    ])
    assert cells_of_rows(html)[0][1] == "One Two"
    assert cells_of_rows(html)[1][1] == ""


def test_rows_without_a_set_height_keep_the_content_in_the_first_row() -> None:
    html = convert([
        row([cell(["Label"]), cell(["First line", "Second line"], "restart", before=240)]),
        row([cell(["Other"]), cell([], "continue")]),
    ])
    assert cells_of_rows(html)[0][1] == "First line Second line"


def test_a_merged_cell_with_one_paragraph_is_untouched() -> None:
    html = convert([
        row([cell(["Label"]), cell(["Only line"], "restart")], height=300),
        row([cell(["Other"]), cell([], "continue")], height=600),
    ])
    assert cells_of_rows(html) == [["Label", "Only line"], ["Other", ""]]


def test_imported_list_items_are_flagged_so_the_editor_drops_list_margins() -> None:
    document = docx.Document()
    document.add_paragraph("A bullet", style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)
    html = convert_docx_bytes_to_html(buffer.getvalue())["html"]
    assert re.search(r'<li data-flat="1"', html)
