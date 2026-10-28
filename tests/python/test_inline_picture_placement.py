"""Tests that a picture alone in its paragraph takes the paragraph's alignment and indent."""
import io
import os
import re
import sys

import cv2
import docx
import numpy as np
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "ml", "src"))

from converter import convert_docx_bytes_to_html  # noqa: E402


def convert(alignment=None, indent_in: float = 0.0, text: str = "") -> str:
    """A DOCX with one 2in picture paragraph (optionally aligned, indented, with text), converted to HTML."""
    document = docx.Document()
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(1)
    paragraph = document.add_paragraph()
    if alignment is not None:
        paragraph.alignment = alignment
    if indent_in:
        paragraph.paragraph_format.left_indent = Inches(indent_in)
    ok, encoded = cv2.imencode(".png", np.full((60, 120, 3), 160, dtype=np.uint8))
    assert ok
    paragraph.add_run().add_picture(io.BytesIO(encoded.tobytes()), width=Inches(2), height=Inches(1))
    if text:
        paragraph.add_run(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return convert_docx_bytes_to_html(buffer.getvalue())["html"]


def placement(html: str) -> tuple[str | None, str | None]:
    """(data-alignment, margin of the container style) of the first picture."""
    tag = re.search(r"<img[^>]*>", html).group(0)
    side = re.search(r'data-alignment="(\w+)"', tag)
    container = re.search(r'containerstyle="[^"]*margin:\s*([^;]+);', tag)
    return (side.group(1) if side else None, container.group(1) if container else None)


def test_picture_in_a_left_aligned_paragraph_sits_at_the_left() -> None:
    assert placement(convert()) == ("left", "0 auto 0 0")  # Word's default alignment is left


def test_picture_in_a_justified_paragraph_sits_at_the_left() -> None:
    assert placement(convert(WD_ALIGN_PARAGRAPH.JUSTIFY)) == ("left", "0 auto 0 0")


def test_picture_in_a_right_aligned_paragraph_sits_at_the_right() -> None:
    assert placement(convert(WD_ALIGN_PARAGRAPH.RIGHT)) == ("right", "0 0 0 auto")


def test_picture_in_a_centered_paragraph_stays_centered() -> None:
    assert placement(convert(WD_ALIGN_PARAGRAPH.CENTER)) == (None, None)


def test_picture_in_an_indented_paragraph_keeps_the_indent() -> None:
    assert placement(convert(indent_in=1.0)) == ("left", "0 auto 0 96.0px")


def test_picture_beside_text_in_the_same_paragraph_is_left_alone() -> None:
    """Word lays the text and the picture out on one line, which a block cannot do."""
    assert placement(convert(text=" caption")) == (None, None)
