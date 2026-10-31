"""Equations in a DOCX reach the editor as math nodes: inline spans, display blocks, numbers, and a text fallback."""
import io
import os
import sys

import docx
from lxml import etree

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "ml", "src"))

from converter import convert_docx_bytes_to_html  # noqa: E402

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = f'xmlns:m="{M}" xmlns:w="{W}"'

SQUARE = '<m:oMath><m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>'
UNKNOWN = "<m:oMath><m:mystery><m:r><m:t>q</m:t></m:r></m:mystery></m:oMath>"


def run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


TAB = "<w:r><w:tab/></w:r>"


def convert(*paragraphs: str, table: bool = False) -> str:
    """HTML for a document whose paragraphs hold the given raw WordprocessingML children."""
    document = docx.Document()
    body = document.element.body
    sect_pr = body.find(f"{{{W}}}sectPr")
    container = body
    if table:
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        container = cell._tc
        for existing in list(container.findall(f"{{{W}}}p")):
            container.remove(existing)
    for children in paragraphs:
        paragraph = etree.fromstring(f"<w:p {NS}>{children}</w:p>")
        if table:
            container.append(paragraph)
        elif sect_pr is not None:
            sect_pr.addprevious(paragraph)
        else:
            body.append(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return convert_docx_bytes_to_html(buffer.getvalue())["html"]


def test_equation_inside_a_sentence_is_an_inline_node() -> None:
    html = convert(run("Where ") + SQUARE + run(" is the area."))
    assert '<span data-type="inline-math" data-latex="x^{2}"></span>' in html
    assert "Where " in html and "is the area." in html
    assert 'data-type="block-math"' not in html


def test_paragraph_that_is_only_an_equation_is_a_display_block() -> None:
    html = convert(SQUARE)
    assert '<div data-type="block-math" data-latex="x^{2}"></div>' in html
    assert "<p" not in html.split('<div data-type="block-math"')[1].split("</div>")[0]


def test_math_paragraph_element_is_a_display_block() -> None:
    html = convert(f"<m:oMathPara>{SQUARE}</m:oMathPara>")
    assert '<div data-type="block-math" data-latex="x^{2}"></div>' in html


def test_number_after_a_tab_belongs_to_the_equation() -> None:
    html = convert(TAB + SQUARE + TAB + run("(7)"))
    assert '<div data-type="block-math" data-latex="x^{2}" data-number="7"></div>' in html
    assert "(7)" not in html.replace('data-number="7"', "")


def test_number_typed_inside_the_equation_is_split_off() -> None:
    typed = '<m:oMath><m:r><m:t>x=1</m:t></m:r><m:r><m:t>#(4)</m:t></m:r></m:oMath>'
    html = convert(typed)
    assert 'data-latex="x=1" data-number="4"' in html


def test_text_that_is_not_a_number_keeps_the_equation_inline() -> None:
    html = convert(SQUARE + run(", see below"))
    assert 'data-type="inline-math"' in html
    assert 'data-type="block-math"' not in html


def test_equation_that_cannot_be_converted_keeps_its_text() -> None:
    html = convert(run("Note: ") + UNKNOWN)
    assert "Note: " in html
    assert html.rstrip().endswith("q</p>")  # the equation's own text follows the sentence
    assert "data-type" not in html


def test_unconvertible_display_equation_is_not_dropped() -> None:
    html = convert(UNKNOWN)
    assert ">q<" in html
    assert 'data-type="block-math"' not in html


def test_display_equation_in_a_table_cell_is_not_wrapped_in_a_paragraph() -> None:
    html = convert(SQUARE, table=True)
    cell = html.split("<td")[1]
    assert '<div data-type="block-math" data-latex="x^{2}"></div>' in cell
    assert '<p style="margin:0' not in cell.split('<div data-type="block-math"')[0].split(">", 1)[1]


def test_latex_with_quotes_is_attribute_safe() -> None:
    quoted = '<m:oMath><m:r><m:rPr><m:nor/></m:rPr><m:t>say "hi" &amp; &lt;go&gt;</m:t></m:r></m:oMath>'
    html = convert(run("x ") + quoted)
    assert 'data-latex="\\text{say &quot;hi&quot; \\&amp; &lt;go&gt;}"' in html
