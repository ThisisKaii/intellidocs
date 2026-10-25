"""Imported headings keep their own spacing, and cover-page lines are not promoted to headings."""
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ML_SRC = Path(__file__).resolve().parents[2] / "ml" / "src"
if str(ML_SRC) not in sys.path:
    sys.path.insert(0, str(ML_SRC))

from converter import _docx_document_to_result, _infer_heading_styles, _page_break_paragraphs  # noqa: E402


def _heading_html(before_pt=None, after_pt=None):
    """A Heading 1 paragraph with optional direct spacing, converted to HTML."""
    doc = docx.Document()
    p = doc.add_paragraph("Introduction", style="Heading 1")
    if before_pt is not None:
        p.paragraph_format.space_before = Pt(before_pt)
    if after_pt is not None:
        p.paragraph_format.space_after = Pt(after_pt)
    return _docx_document_to_result(doc)["html"]


class TestHeadingSpacing:
    def test_direct_spacing_is_emitted_instead_of_the_fixed_default(self):
        html = _heading_html(before_pt=36, after_pt=0)
        assert "margin-top:48.0px" in html or "margin-top:48px" in html
        assert "margin-bottom:0" in html
        assert "0.85em" not in html

    def test_spacing_inherited_from_the_style_is_used_too(self):
        # python-docx's Heading 1 style defines space-before, so no fixed margin is needed.
        html = _heading_html()
        assert "<h1" in html and "0.85em" not in html

    def test_size_comes_from_the_document_not_from_a_constant(self):
        doc = docx.Document()
        doc.styles["Heading 1"].font.size = Pt(13)
        doc.add_paragraph("Methods", style="Heading 1")
        html = _docx_document_to_result(doc)["html"]
        assert "font-size:13pt" in html and "font-size:17pt" not in html

    def test_alignment_and_weight_are_still_present(self):
        doc = docx.Document()
        p = doc.add_paragraph("Chapter One", style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        html = _docx_document_to_result(doc)["html"]
        assert "text-align:center" in html and "font-weight:bold" in html


def _styled(doc, text, size, bold=True, center=True):
    """Add a paragraph whose one run has an explicit size, like pdf2docx output."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _thesis_like_doc():
    doc = docx.Document()
    cover = [_styled(doc, "INTELLIDOCS: AN INTELLIGENT DOCUMENT EDITING SYSTEM", 17),
             _styled(doc, "Presented to the Faculty of the College of Computer Studies", 17, bold=False),
             _styled(doc, "CHAPTER I", 17)]
    doc.add_page_break()
    later = _styled(doc, "Background of the Study", 17, center=False)
    for _ in range(12):
        body = doc.add_paragraph()
        body.add_run("Plain body text that is long enough to set the dominant size of the document. " * 4).font.size = Pt(11)
    return doc, cover, later


class TestCoverPage:
    def test_centred_cover_lines_are_not_promoted(self):
        doc, cover, _ = _thesis_like_doc()
        _infer_heading_styles(doc, _page_break_paragraphs(doc))
        title, presented = (docx.text.paragraph.Paragraph(c._p, c._parent) for c in cover[:2])
        assert title.style.name == "Normal"
        assert presented.style.name == "Normal"

    def test_a_chapter_line_on_the_cover_page_is_still_a_heading(self):
        doc, cover, _ = _thesis_like_doc()
        _infer_heading_styles(doc, _page_break_paragraphs(doc))
        assert cover[2].style.name.startswith("Heading")

    def test_headings_after_the_first_page_are_promoted(self):
        doc, _, later = _thesis_like_doc()
        _infer_heading_styles(doc, _page_break_paragraphs(doc))
        assert later.style.name.startswith("Heading")

    def test_without_page_info_behaviour_is_unchanged(self):
        doc, cover, _ = _thesis_like_doc()
        _infer_heading_styles(doc)
        assert cover[0].style.name.startswith("Heading")

    def test_cover_lines_keep_their_direct_size_in_the_html(self):
        doc, _, _ = _thesis_like_doc()
        pb = _page_break_paragraphs(doc)
        _infer_heading_styles(doc, pb)
        html = _docx_document_to_result(doc, pb)["html"]
        first = re.search(r"<(h1|p)[^>]*>(?:(?!</).)*INTELLIDOCS", html, re.S)
        assert first and first.group(1) == "p"
