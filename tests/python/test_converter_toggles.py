"""Bold and italic are toggle properties: a character style and a paragraph style that both set them cancel."""
import io
import os
import sys
from typing import Optional

import docx
from docx.enum.style import WD_STYLE_TYPE

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "ml", "src"))

from converter import convert_docx_bytes_to_html  # noqa: E402


def html_of(paragraph_bold: Optional[bool], character_bold: Optional[bool], direct_bold: Optional[bool] = None, italic: bool = False) -> str:
    """HTML of one paragraph with one run, given what the paragraph style, character style and run say about bold."""
    document = docx.Document()
    paragraph_style = document.styles.add_style("TestPara", WD_STYLE_TYPE.PARAGRAPH)
    character_style = document.styles.add_style("TestChar", WD_STYLE_TYPE.CHARACTER)
    if italic:
        paragraph_style.font.italic = paragraph_bold
        character_style.font.italic = character_bold
    else:
        paragraph_style.font.bold = paragraph_bold
        character_style.font.bold = character_bold
    paragraph = document.add_paragraph(style="TestPara")
    run = paragraph.add_run("sample")
    run.style = "TestChar"
    if direct_bold is not None:
        if italic:
            run.font.italic = direct_bold
        else:
            run.font.bold = direct_bold
    buffer = io.BytesIO()
    document.save(buffer)
    return convert_docx_bytes_to_html(buffer.getvalue())["html"]


def is_bold(html: str) -> bool:
    return "<strong>" in html


def test_a_bold_paragraph_style_alone_is_bold() -> None:
    assert is_bold(html_of(True, None))


def test_a_bold_character_style_alone_is_bold() -> None:
    assert is_bold(html_of(None, True))


def test_both_styles_bold_cancel_out() -> None:
    assert not is_bold(html_of(True, True))


def test_a_character_style_that_turns_bold_off_in_a_bold_paragraph_does_not_double_toggle() -> None:
    # The character style says "not bold" explicitly; only one style is on, so the text is bold.
    assert is_bold(html_of(True, False))


def test_a_direct_setting_on_the_run_is_absolute() -> None:
    assert is_bold(html_of(True, True, direct_bold=True))
    assert not is_bold(html_of(True, None, direct_bold=False))


def test_italic_toggles_the_same_way() -> None:
    assert "<em>" in html_of(True, None, italic=True)
    assert "<em>" not in html_of(True, True, italic=True)


def test_nothing_set_anywhere_is_plain() -> None:
    assert not is_bold(html_of(None, None))
