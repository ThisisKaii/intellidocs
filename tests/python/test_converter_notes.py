"""Footnotes and endnotes: a raised number in the text, the note's text collected at the end in reference order."""
import io
import os
import re
import sys
from typing import Dict, List

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "ml", "src"))

from converter import convert_docx_bytes_to_html  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def add_notes_part(document: docx.Document, kind: str, notes: Dict[int, str]) -> None:
    """Attach a footnotes.xml or endnotes.xml holding the given notes (id -> text), plus Word's separators."""
    paragraphs = "".join(
        f'<w:{kind} w:id="{note_id}"><w:p><w:r><w:{kind}Ref/></w:r><w:r><w:t xml:space="preserve"> {text}</w:t></w:r></w:p></w:{kind}>'
        for note_id, text in notes.items()
    )
    xml = (
        f'<w:{kind}s xmlns:w="{W}"><w:{kind} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{kind}>'
        f'<w:{kind} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{kind}>{paragraphs}</w:{kind}s>'
    )
    part = Part(
        PackURI(f"/word/{kind}s.xml"),
        f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
        xml.encode("utf-8"),
        document.part.package,
    )
    document.part.relate_to(part, RT.FOOTNOTES if kind == "footnote" else RT.ENDNOTES)


def reference(kind: str, note_id: int) -> str:
    return f'<w:r xmlns:w="{W}"><w:rPr><w:rStyle w:val="{kind.capitalize()}Reference"/></w:rPr><w:{kind}Reference w:id="{note_id}"/></w:r>'


def text_run(text: str) -> str:
    return f'<w:r xmlns:w="{W}"><w:t xml:space="preserve">{text}</w:t></w:r>'


def convert(build) -> str:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return convert_docx_bytes_to_html(buffer.getvalue())["html"]


def paragraph(document: docx.Document, *runs: str):
    p = document.add_paragraph()
    for run in runs:
        p._p.append(etree.fromstring(run))
    return p


def sups(html: str) -> List[str]:
    return re.findall(r"<sup>([^<]*)</sup>", html)


def test_footnotes_are_numbered_in_reference_order_not_by_id() -> None:
    def build(document: docx.Document) -> None:
        paragraph(document, text_run("First"), reference("footnote", 5), text_run(" then second"), reference("footnote", 3))
        add_notes_part(document, "footnote", {3: "Note three.", 5: "Note five."})

    html = convert(build)
    assert sups(html) == ["1", "2", "1", "2"]  # the two references, then the two notes
    notes = html.split("Notes</p>")[1]
    assert notes.index("Note five.") < notes.index("Note three.")  # id 5 was referenced first
    assert re.search(r"<sup>1</sup> <span[^>]*>\s*Note five\.", notes)


def test_a_note_is_collected_once_however_many_times_its_text_is_walked() -> None:
    def build(document: docx.Document) -> None:
        paragraph(document, text_run("Once"), reference("footnote", 2))
        add_notes_part(document, "footnote", {2: "Only note."})

    html = convert(build)
    assert html.count("Only note.") == 1


def test_endnotes_are_numbered_in_roman_under_their_own_label() -> None:
    def build(document: docx.Document) -> None:
        paragraph(document, text_run("Text"), reference("endnote", 1), text_run(" more"), reference("endnote", 2))
        add_notes_part(document, "endnote", {1: "First end.", 2: "Second end."})

    html = convert(build)
    assert sups(html) == ["i", "ii", "i", "ii"]
    assert "Endnotes</p>" in html and "Notes</p>" not in html.replace("Endnotes</p>", "")


def test_a_reference_in_a_table_cell_is_counted() -> None:
    def build(document: docx.Document) -> None:
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0]._p.append(etree.fromstring(text_run("In a cell")))
        cell.paragraphs[0]._p.append(etree.fromstring(reference("footnote", 1)))
        add_notes_part(document, "footnote", {1: "Cell note."})

    html = convert(build)
    assert sups(html) == ["1", "1"]
    assert "Cell note." in html


def test_a_document_without_notes_gets_no_notes_section() -> None:
    document = docx.Document()
    document.add_paragraph("Plain text.")
    buffer = io.BytesIO()
    document.save(buffer)
    html = convert_docx_bytes_to_html(buffer.getvalue())["html"]
    assert "Notes" not in html and "<sup>" not in html


def test_a_reference_to_a_missing_note_still_shows_its_number() -> None:
    def build(document: docx.Document) -> None:
        paragraph(document, text_run("Dangling"), reference("footnote", 9))

    html = convert(build)
    assert sups(html) == ["1"]
