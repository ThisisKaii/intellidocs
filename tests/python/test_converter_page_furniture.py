"""Headers, footers, page-pinned pictures, sections and PDF page furniture survive an import."""
import io
import sys
from pathlib import Path

import docx
import pymupdf
import pytest
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
ML_SRC = ROOT / "ml" / "src"
if str(ML_SRC) not in sys.path:
    sys.path.insert(0, str(ML_SRC))

from converter import (  # noqa: E402
    _docx_document_to_result,
    _furniture_to_layout,
    _printed_page_number,
    _realign_from_pdf,
)

WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def _png(width: int = 40, height: int = 20) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    pix.set_rect(pix.irect, (30, 90, 200))
    return pix.tobytes("png")


def _anchor(document: docx.Document, x_in: float, y_in: float, relative: str = "page") -> None:
    """Add a picture and turn it into one Word pins to the page at (x, y) inches."""
    document.add_picture(io.BytesIO(_png()), width=Inches(1), height=Inches(0.5))
    inline = document.element.body.find(".//" + f"{{{WP}}}inline")
    anchor = etree.Element(f"{{{WP}}}anchor", behindDoc="0")
    for name, offset in (("positionH", x_in), ("positionV", y_in)):
        position = etree.SubElement(anchor, f"{{{WP}}}{name}", relativeFrom=relative)
        etree.SubElement(position, f"{{{WP}}}posOffset").text = str(int(offset * 914400))
    for child in list(inline):
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


class TestHeadersAndFooters:
    def test_a_header_picture_comes_from_the_headers_own_relationships(self):
        document = docx.Document()
        header = document.sections[0].header
        header.is_linked_to_previous = False
        header.paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(1))
        result = _docx_document_to_result(document)
        parts = result["hf"]["parts"]
        assert len(parts) == 1
        assert "<img" in next(iter(parts.values()))["flow"]
        assert result["hf"]["sections"][0]["header"] in parts

    def test_a_page_field_becomes_the_editors_page_token(self):
        document = docx.Document()
        header = document.sections[0].header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.add_run("Chapter\t")
        field = etree.SubElement(paragraph._p, qn("w:fldSimple"))
        field.set(qn("w:instr"), " PAGE ")
        run = etree.SubElement(field, qn("w:r"))
        etree.SubElement(run, qn("w:t")).text = "7"
        result = _docx_document_to_result(document)
        flow = next(iter(result["hf"]["parts"].values()))["flow"]
        assert "{page}" in flow and ">7<" not in flow

    def test_a_plain_text_header_stays_in_the_old_text_fields_only(self):
        document = docx.Document()
        header = document.sections[0].header
        header.is_linked_to_previous = False
        header.paragraphs[0].text = "Running title"
        result = _docx_document_to_result(document)
        assert result["hf"]["parts"] == {}
        assert result["header"] == "Running title"

    def test_sections_carry_their_numbering_and_first_page_choice(self):
        document = docx.Document()
        document.add_paragraph("front")
        document.add_section()
        document.add_paragraph("body")
        first, second = document.sections
        first.different_first_page_header_footer = True
        number = etree.SubElement(second._sectPr, qn("w:pgNumType"))
        number.set(qn("w:start"), "1")
        number.set(qn("w:fmt"), "lowerRoman")
        sections = _docx_document_to_result(document)["hf"]["sections"]
        assert sections[0]["title_pg"] is True
        assert (sections[1]["number_start"], sections[1]["number_format"]) == (1, "lowerRoman")

    def test_a_page_that_opens_a_section_carries_its_number_on_the_break(self):
        document = docx.Document()
        document.add_paragraph("front")
        document.add_section()
        document.add_paragraph("body")
        html = _docx_document_to_result(document)["html"]
        assert 'data-page-break="true" data-section="1"' in html


class TestPagePinnedPictures:
    def test_a_picture_repeated_at_the_same_spot_on_several_pages_is_running_art(self):
        document = docx.Document()
        for page in range(3):
            _anchor(document, 0.5, 10.0)
            if page < 2:
                document.add_page_break()
        result = _docx_document_to_result(document)
        assert result["page_art"]["repeat"].count("<img") == 1
        assert result["page_art"]["pages"] == {}
        assert "position:absolute" in result["page_art"]["repeat"]
        assert "<img" not in result["html"]

    def test_a_picture_on_one_page_belongs_to_that_page(self):
        document = docx.Document()
        document.add_paragraph("one")
        document.add_page_break()
        _anchor(document, 1.0, 2.0)
        result = _docx_document_to_result(document)
        assert list(result["page_art"]["pages"]) == ["2"]
        assert "left:96px" in result["page_art"]["pages"]["2"] and "top:192px" in result["page_art"]["pages"]["2"]

    def test_a_picture_that_follows_its_paragraph_stays_in_the_text(self):
        document = docx.Document()
        _anchor(document, 1.0, 1.0, relative="paragraph")
        result = _docx_document_to_result(document)
        assert "<img" in result["html"]
        assert result["page_art"] == {"repeat": "", "pages": {}, "overlays": {}}


class TestTableRows:
    def test_row_height_is_a_minimum_not_a_fixed_height(self):
        document = docx.Document()
        table = document.add_table(rows=1, cols=1)
        row_properties = table.rows[0]._tr.get_or_add_trPr()
        etree.SubElement(row_properties, qn("w:trHeight")).set(qn("w:val"), "1440")
        html = _docx_document_to_result(document)["html"]
        assert "min-height:96px" in html and "height:96px" not in html.replace("min-height:96px", "")


class TestPdfRealign:
    def _pdf(self, text: str, x0: float) -> bytes:
        pdf = pymupdf.open()
        page = pdf.new_page(width=612, height=792)
        page.insert_text((x0, 100), text, fontsize=12)
        data = pdf.tobytes()
        pdf.close()
        return data

    def _document(self, text: str, alignment: str) -> docx.Document:
        document = docx.Document()
        section = document.sections[0]
        section.left_margin = section.right_margin = Inches(1)
        paragraph = document.add_paragraph(text)
        paragraph.alignment = {"right": 2, "center": 1, "left": 0}[alignment]
        return document

    def test_a_title_the_pdf_drew_centred_is_centred_again(self):
        width = pymupdf.get_text_length("LIST OF TABLES", fontsize=12)
        pdf = self._pdf("LIST OF TABLES", (612 - width) / 2)
        document = self._document("LIST OF TABLES", "right")
        assert _realign_from_pdf(document, pdf, set()) == 1
        assert document.paragraphs[0]._p.find(qn("w:pPr") + "/" + qn("w:jc")).get(qn("w:val")) == "center"

    def test_a_line_drawn_at_the_left_margin_is_left_aligned(self):
        pdf = self._pdf("Heading", 72)
        document = self._document("Heading", "center")
        _realign_from_pdf(document, pdf, set())
        assert document.paragraphs[0]._p.find(qn("w:pPr") + "/" + qn("w:jc")).get(qn("w:val")) == "left"


class TestPdfFurniture:
    @pytest.mark.parametrize(
        "text, expected",
        [("12", (12, "decimal")), ("iv", (4, "lowerRoman")), ("XIV", (14, "upperRoman")), ("Page 3 of 9", (3, "decimal")), ("Chapter", None)],
    )
    def test_printed_page_numbers_are_read(self, text, expected):
        assert _printed_page_number(text) == expected

    def test_stripped_page_numbers_return_as_a_header_with_their_own_numbering(self):
        sections = [
            {"index": index, "header": None, "footer": None, "first_header": None, "first_footer": None, "title_pg": False,
             "number_start": None, "number_format": "decimal", "header_distance": 48, "footer_distance": 48}
            for index in range(4)
        ]
        result = {"html": '<p>a</p><hr class="page-break" data-section="1" /><hr class="page-break" data-section="2" /><hr class="page-break" data-section="3" />',
                  "hf": {"parts": {}, "sections": sections}}
        report = {"items": [
            {"page": 1, "side": "header", "key": "<page-number>|", "text": "i", "jc": "right", "has_image": False},
            {"page": 2, "side": "header", "key": "<page-number>|", "text": "ii", "jc": "right", "has_image": False},
            {"page": 3, "side": "header", "key": "<page-number>|", "text": "1", "jc": "right", "has_image": False},
        ]}
        _furniture_to_layout(result, report, 4)
        assert "{page}" in result["hf"]["parts"]["pdf_header"]["flow"]
        assert [section["header"] for section in result["hf"]["sections"]] == [None, "pdf_header", "pdf_header", "pdf_header"]
        assert [(s["number_start"], s["number_format"]) for s in result["hf"]["sections"][1:]] == [(1, "lowerRoman"), (None, "lowerRoman"), (1, "decimal")]


def _hang_box(document: docx.Document, paragraph: "docx.text.paragraph.Paragraph", text: str) -> None:
    """Hang a no-wrap text box beside a paragraph, 1in from the page's left edge and 0.5in down from the paragraph."""
    wps = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    wp = WP
    xml = (
        f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:drawing>'
        f'<wp:anchor xmlns:wp="{wp}" behindDoc="0"><wp:extent cx="1828800" cy="914400"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>914400</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>457200</wp:posOffset></wp:positionV><wp:wrapNone/>'
        f'<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData>'
        f'<wps:wsp xmlns:wps="{wps}"><wps:spPr><a:solidFill><a:srgbClr val="DDEEFF"/></a:solidFill></wps:spPr>'
        f'<wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx></wps:wsp>'
        f'</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'
    )
    paragraph._p.append(etree.fromstring(xml))


class TestBoxesBesideParagraphs:
    def test_a_text_box_hung_beside_a_paragraph_becomes_an_overlay_tagged_on_that_paragraph(self):
        document = docx.Document()
        paragraph = document.add_paragraph("Body text")
        _hang_box(document, paragraph, "Abstract box")
        result = _docx_document_to_result(document)
        overlays = result["page_art"]["overlays"]
        assert len(overlays) == 1
        box = next(iter(overlays.values()))
        assert (box["x"], box["y"], box["w"], box["h"]) == (96, 48, 192, 96)
        assert "Abstract box" in box["html"] and "background:#DDEEFF" in box["html"]
        assert 'data-overlay="1"' in result["html"]
        assert "Abstract box" not in result["html"]

    def test_a_small_badge_inside_a_line_of_text_is_dropped(self):
        document = docx.Document()
        paragraph = document.add_paragraph("A. Author ")
        paragraph.add_run().add_picture(io.BytesIO(_png()), width=Inches(0.17))
        assert "<img" not in _docx_document_to_result(document)["html"]

    def test_a_picture_alone_in_its_paragraph_stays(self):
        document = docx.Document()
        document.add_picture(io.BytesIO(_png()), width=Inches(0.17))
        assert "<img" in _docx_document_to_result(document)["html"]


class TestPageSetupAndFirstSection:
    def test_the_margins_most_sections_share_win_over_a_cover_section(self):
        document = docx.Document()
        document.sections[0].top_margin = Inches(0.25)
        for _ in range(2):
            section = document.add_section()
            section.top_margin = Inches(1.0)
        assert _docx_document_to_result(document)["page_setup"]["margins"]["top"] == 1.0

    def test_the_cover_section_keeps_its_own_top_margin_by_shifting_its_first_block(self):
        document = docx.Document()
        document.sections[0].top_margin = Inches(0.25)
        document.add_paragraph("Cover title")
        for _ in range(2):
            document.add_section().top_margin = Inches(1.0)
            document.add_paragraph("Body")
        html = _docx_document_to_result(document)["html"]
        assert "margin-top:-72px" in html.split("Body")[0]
