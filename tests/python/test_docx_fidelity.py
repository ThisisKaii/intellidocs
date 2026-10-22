"""Imported DOCX fidelity: explicit indents/sizes, Word-accurate line spacing, header page fields."""
from pathlib import Path
import re
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parents[2]
ML_SRC = REPO_ROOT / "ml" / "src"
if str(ML_SRC) not in sys.path:
    sys.path.insert(0, str(ML_SRC))

from converter import (  # noqa: E402
    _docx_document_to_result,
    _extract_header_footer,
    _paragraph_text_and_page_format,
)


def _body_html(doc: docx.Document) -> str:
    return _docx_document_to_result(doc)["html"]


def _add_field(paragraph: docx.text.paragraph.Paragraph, instruction: str, cached: str) -> None:
    """Append a complex field (begin / instrText / separate / cached result / end) to a paragraph."""
    def run_with(child: object) -> None:
        run = paragraph.add_run()
        run._r.append(child)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    instr = OxmlElement("w:instrText")
    instr.text = f" {instruction} "
    run_with(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run_with(separate)
    paragraph.add_run(cached)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


class TestParagraphIndentAndSize:
    """Explicit text-indent / font-size keep an applied preset from rewriting imported text."""

    def test_unindented_paragraph_gets_explicit_zero_indent(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("A Proposal Project")
        assert "text-indent:0" in _body_html(doc)

    def test_word_first_line_indent_is_kept(self) -> None:
        doc = docx.Document()
        paragraph = doc.add_paragraph("Body text")
        paragraph.paragraph_format.first_line_indent = Pt(36)  # 0.5in = 720 twips
        assert "text-indent:48.0px" in _body_html(doc)

    def test_style_defined_indent_resolves_through_the_style_chain(self) -> None:
        doc = docx.Document()
        style = doc.styles["Normal"]
        style.paragraph_format.first_line_indent = Pt(18)
        doc.add_paragraph("Inherits an indent from its style")
        assert "text-indent:24.0px" in _body_html(doc)

    def test_base_font_size_is_explicit(self) -> None:
        doc = docx.Document()
        doc.styles["Normal"].font.size = Pt(11)
        doc.add_paragraph("Sized by style")
        assert re.search(r'<p style="[^"]*font-size:11pt', _body_html(doc))

    def test_table_cell_paragraphs_are_not_indented_by_a_preset(self) -> None:
        doc = docx.Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "cell"
        assert "text-indent:0" in _body_html(doc)


class TestAutoLineSpacing:
    """Word's 'auto' spacing multiplies the font's natural line height, not 1em."""

    def _line_height(self, font_name: str) -> float:
        doc = docx.Document()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.add_run("Spaced").font.name = font_name
        return float(re.search(r"line-height:([\d.]+)", _body_html(doc)).group(1))

    def test_times_new_roman_one_and_a_half(self) -> None:
        assert self._line_height("Times New Roman") == 1.725

    def test_calibri_has_taller_natural_lines(self) -> None:
        assert self._line_height("Calibri") == 1.832

    def test_unknown_font_uses_the_default_factor(self) -> None:
        assert self._line_height("Some Unknown Face") == 1.725

    def test_blank_paragraph_uses_its_paragraph_mark_font(self) -> None:
        doc = docx.Document()
        doc.styles["Normal"].font.name = "Cambria"
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.5
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Times New Roman")
        r_pr = OxmlElement("w:rPr")
        r_pr.append(r_fonts)
        paragraph._p.get_or_add_pPr().append(r_pr)
        assert "line-height:1.725" in _body_html(doc)


class TestHeaderPageFields:
    """A PAGE field becomes a page-number format; its frozen cached number is dropped."""

    def test_complex_page_field_reports_number_and_drops_cached_text(self) -> None:
        doc = docx.Document()
        _add_field(doc.sections[0].header.paragraphs[0], "PAGE", "2")
        header, _footer, header_format, footer_format = _extract_header_footer(doc)
        assert (header, header_format, footer_format) == ("", "number", "none")

    def test_roman_switch_reports_roman(self) -> None:
        doc = docx.Document()
        _add_field(doc.sections[0].footer.paragraphs[0], r"PAGE \* roman", "ii")
        _header, footer, _hf, footer_format = _extract_header_footer(doc)
        assert (footer, footer_format) == ("", "roman")

    def test_text_around_the_field_is_kept(self) -> None:
        doc = docx.Document()
        paragraph = doc.sections[0].header.paragraphs[0]
        paragraph.add_run("Running Title ")
        _add_field(paragraph, "PAGE", "7")
        header, _footer, header_format, _ff = _extract_header_footer(doc)
        assert (header, header_format) == ("Running Title", "number")

    def test_other_fields_keep_their_result_text(self) -> None:
        doc = docx.Document()
        paragraph = doc.sections[0].header.paragraphs[0]
        _add_field(paragraph, "NUMPAGES", "33")
        text, page_format = _paragraph_text_and_page_format(paragraph._element)
        assert (text, page_format) == ("33", "none")

    def test_document_without_header_fields_is_unchanged(self) -> None:
        doc = docx.Document()
        assert _extract_header_footer(doc) == ("", "", "none", "none")
