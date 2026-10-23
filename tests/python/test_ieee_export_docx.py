"""IEEE Word export: page setup, fonts and the two-column body of a generated sample."""
from pathlib import Path
import os
import shutil
import subprocess

import docx
import pytest
from docx.oxml.ns import qn
from docx.shared import Inches

REPO_ROOT = Path(__file__).resolve().parents[2]
FRONTEND = REPO_ROOT / "frontend"


@pytest.fixture(scope="module")
def sample(tmp_path_factory: pytest.TempPathFactory) -> docx.document.Document:
    """Generate a sample with the frontend vitest suite (which writes IEEE_SAMPLE_OUT)."""
    npx = shutil.which("npx")
    if npx is None or not (FRONTEND / "node_modules" / "docx").exists():
        pytest.skip("frontend dependencies are not installed")
    out = tmp_path_factory.mktemp("ieee") / "sample.docx"
    env = {**os.environ, "IEEE_SAMPLE_OUT": str(out)}
    subprocess.run(
        [npx, "vitest", "run", "src/lib/__tests__/ieeeExport.test.ts"],
        cwd=FRONTEND, env=env, check=True, capture_output=True,
    )
    return docx.Document(str(out))


def test_two_sections_second_is_continuous_two_column(sample: docx.document.Document) -> None:
    """Front matter is full width; the body is a continuous 2-column section with a 0.25in gap."""
    assert len(sample.sections) == 2
    first, second = sample.sections
    assert first._sectPr.find(qn("w:cols")) is None
    assert second.start_type.name == "CONTINUOUS"
    cols = second._sectPr.find(qn("w:cols"))
    assert cols.get(qn("w:num")) == "2"
    assert int(cols.get(qn("w:space"))) == 360


def test_letter_page_and_margins(sample: docx.document.Document) -> None:
    """Both sections are US Letter with the IEEE margins."""
    for section in sample.sections:
        assert section.page_width == Inches(8.5)
        assert section.page_height == Inches(11)
        assert section.top_margin == Inches(0.75)
        assert section.bottom_margin == Inches(1)
        assert section.left_margin == Inches(0.625)
        assert section.right_margin == Inches(0.625)


def test_fonts_and_sizes(sample: docx.document.Document) -> None:
    """Times New Roman throughout; 24pt title, 9pt bold abstract, 9pt hanging references."""
    paragraphs = {p.text.split("—")[0].strip(): p for p in sample.paragraphs if p.text.strip()}
    title = paragraphs["Deep Learning for Slide Recognition"]
    assert title.runs[0].font.size.pt == 24
    assert title.alignment.name == "CENTER"
    abstract = paragraphs["Abstract"]
    assert all(r.font.size.pt == 9 and r.bold for r in abstract.runs)
    ref = next(p for p in sample.paragraphs if p.text.startswith("[1]"))
    assert ref.runs[0].font.size.pt == 9
    assert ref.paragraph_format.first_line_indent < 0
    heading = paragraphs["I. Introduction"]
    assert heading.alignment.name == "CENTER"
    assert heading.runs[0].font.small_caps
    assert all(r.font.name == "Times New Roman" for p in sample.paragraphs for r in p.runs)
