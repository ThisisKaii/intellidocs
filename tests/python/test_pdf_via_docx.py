"""PDF imports are normalized to DOCX (pdf2docx), repaired, and share the DOCX converter."""
from pathlib import Path
import io
import re

import pytest
import sys

import docx
from docx.enum.section import WD_SECTION
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parents[2]
ML_SRC = REPO_ROOT / "ml" / "src"
if str(ML_SRC) not in sys.path:
    sys.path.insert(0, str(ML_SRC))

from converter import (  # noqa: E402
    _docx_document_to_result,
    _heading_candidate,
    _infer_heading_styles,
    _page_break_paragraphs,
    _strip_running_furniture,
    convert_pdf_bytes_to_html,
)

FIXTURE = REPO_ROOT / "tests" / "fixtures" / "import_baselines" / "documents" / "V2 INTELLIDOCS sabotage.pdf"


def _paged_document(pages: int = 5) -> docx.Document:
    """Build a DOCX with a running header and page number on every page."""
    doc = docx.Document()
    for index in range(pages):
        doc.add_paragraph("Running Title of the Study")
        doc.add_paragraph(f"Unique body text for page number {index}")
        doc.add_paragraph(f"{index + 1}")
        doc.add_section(WD_SECTION.NEW_PAGE)
    return doc


def _body_text(doc: docx.Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


class TestPageBreaks:
    """Next-page section breaks become page breaks; continuous ones do not."""

    def test_next_page_sections_emit_one_break_per_page_boundary(self) -> None:
        doc = _paged_document(4)
        html = _docx_document_to_result(doc, _page_break_paragraphs(doc))["html"]
        assert html.count("data-page-break") == 3

    def test_continuous_section_is_not_a_page_break(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("one")
        doc.add_section(WD_SECTION.CONTINUOUS)
        doc.add_paragraph("two")
        assert _page_break_paragraphs(doc) == set()

    def test_docx_next_page_sections_become_page_breaks(self) -> None:
        doc = _paged_document(3)
        assert _docx_document_to_result(doc)["html"].count("data-page-break") == 2


class TestRunningFurniture:
    """Text repeated at page edges is removed; unique body text is kept."""

    def test_repeated_header_and_page_numbers_are_removed(self) -> None:
        doc = _paged_document(5)
        removed = _strip_running_furniture(doc, _page_break_paragraphs(doc))
        text = _body_text(doc)
        assert removed >= 10
        assert "Running Title" not in text
        assert "Unique body text for page number 3" in text

    def test_lines_differing_only_by_a_number_are_kept(self) -> None:
        doc = docx.Document()
        for index in range(5):
            doc.add_paragraph(f"Article {index + 1}")
            doc.add_paragraph(f"Provision text number {index}")
            doc.add_section(WD_SECTION.NEW_PAGE)
        _strip_running_furniture(doc, _page_break_paragraphs(doc))
        text = _body_text(doc)
        assert "Article 3" in text and "Provision text number 3" in text

    def test_short_documents_are_left_alone(self) -> None:
        doc = _paged_document(2)
        assert _strip_running_furniture(doc, _page_break_paragraphs(doc)) == 0

    def test_page_break_survives_when_the_break_paragraph_is_furniture(self) -> None:
        doc = _paged_document(5)
        breaks = _page_break_paragraphs(doc)
        _strip_running_furniture(doc, breaks)
        assert _docx_document_to_result(doc, breaks)["html"].count("data-page-break") == 4


class TestHeadingInference:
    """Heading shape rules and Heading N style assignment."""

    def test_numbered_bold_line_is_a_heading(self) -> None:
        assert _heading_candidate("2.1 Literature Review", 11.0, True, 11.0)

    def test_table_caption_is_not_a_heading(self) -> None:
        assert not _heading_candidate("Table 2", 11.0, True, 11.0)

    def test_tab_leader_toc_line_is_not_a_heading(self) -> None:
        assert not _heading_candidate("REFERENCES \t105", 11.0, True, 11.0)

    def test_unbalanced_fragment_is_not_a_heading(self) -> None:
        assert not _heading_candidate("Coefficient)", 11.0, True, 11.0)

    def test_bold_sentence_is_not_a_heading(self) -> None:
        assert not _heading_candidate("this is a bold wrapped line of body text", 11.0, True, 11.0)

    def test_large_text_gets_a_heading_style_and_body_does_not(self) -> None:
        doc = docx.Document()
        heading = doc.add_paragraph()
        heading.add_run("Introduction").font.size = Pt(16)
        body = doc.add_paragraph()
        body.add_run("A long ordinary sentence of body text that goes on for a while.").font.size = Pt(11)
        assert _infer_heading_styles(doc) == 1
        assert heading.style.name == "Heading 1"
        assert body.style.name != "Heading 1"


@pytest.mark.skipif(not FIXTURE.exists(), reason="import baseline fixtures are not checked in")
class TestFixtureConversion:
    """End to end on a real fixture: the four pdf2docx gaps stay closed."""

    def test_v2_fixture_is_repaired(self) -> None:
        result = convert_pdf_bytes_to_html(FIXTURE.read_bytes())
        html = result["html"]
        assert set(result) == {
            "html", "page_setup", "header", "footer", "header_number_format", "footer_number_format",
        }
        assert result["page_setup"]["page_size"] and result["page_setup"]["margins"]
        assert "<p><br/></p>" not in html
        assert html.count("data-page-break") >= 25
        assert len(re.findall(r"<h[1-3][ >]", html)) >= 5
        assert html.count("<img") <= 12
        assert len(html) < 1_500_000
