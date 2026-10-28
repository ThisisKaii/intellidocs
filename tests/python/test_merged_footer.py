"""A PDF page whose footer lines were merged into one element is still recognised as running footer."""
import sys
from pathlib import Path

import docx
from docx.enum.section import WD_SECTION

ML_SRC = Path(__file__).resolve().parents[2] / "ml" / "src"
if str(ML_SRC) not in sys.path:
    sys.path.insert(0, str(ML_SRC))

from converter import _page_break_paragraphs, _strip_running_furniture  # noqa: E402

ADDRESS = "Room 102, Rizal Building, DepEd Complex, Pasig City 1600"
PHONE = "Telephone Nos.: (+632) 86337206, (+632) 86318494 Email Address: usec.hrod@deped.gov.ph"


def doc_ref(page: int) -> str:
    return f"Doc. Ref. Code DM-OUHROD Rev 00 Page {page} of 24"


def build(merged_page: int, merged_text: str | None = None) -> docx.document.Document:
    """Six pages; each has a body line and three footer lines, except `merged_page` which has one merged footer."""
    document = docx.Document()
    for page in range(1, 7):
        document.add_paragraph(f"Body of page {page} discussing topic number {page} of the report.")
        if page == merged_page:
            # pdf2docx interleaved the three footer lines into one element
            document.add_paragraph(merged_text or f"{ADDRESS} {doc_ref(page)} {PHONE}")
        else:
            document.add_paragraph(ADDRESS)
            document.add_paragraph(PHONE)
            document.add_paragraph(doc_ref(page))
        if page < 6:
            document.add_section(WD_SECTION.NEW_PAGE)
    return document


def remaining_text(document: docx.document.Document) -> str:
    return " ".join(paragraph.text for paragraph in document.paragraphs)


def test_a_footer_merged_into_one_element_is_stripped_like_the_others() -> None:
    document = build(merged_page=4)
    _strip_running_furniture(document, _page_break_paragraphs(document))
    text = remaining_text(document)
    assert "Rizal" not in text and "Doc. Ref." not in text and "Telephone" not in text
    assert all(f"Body of page {page}" in text for page in range(1, 7))


def test_a_body_paragraph_that_only_mentions_the_footer_text_is_kept() -> None:
    """The footer line is a small part of a real paragraph, so the paragraph stays."""
    sentence = f"{ADDRESS} is where the review committee meets on the first Monday of every month to go over submissions."
    document = build(merged_page=4, merged_text=sentence)
    _strip_running_furniture(document, _page_break_paragraphs(document))
    assert "review committee meets" in remaining_text(document)


def test_a_short_body_line_is_never_taken_for_a_merged_footer() -> None:
    document = build(merged_page=4, merged_text="Pasig City 1600 office")
    _strip_running_furniture(document, _page_break_paragraphs(document))
    assert "Pasig City 1600 office" in remaining_text(document)


def test_a_footer_on_a_short_page_is_not_also_a_header() -> None:
    """On a page with almost nothing but its footer, the footer sits in the top half; the other pages decide its side."""
    document = docx.Document()
    for page in range(1, 6):
        if page != 3:  # page 3 holds only footer lines
            document.add_paragraph(f"Body of page {page} discussing topic number {page} of the report.")
        document.add_paragraph(ADDRESS)
        document.add_paragraph(PHONE)
        document.add_paragraph(doc_ref(page))
        if page < 5:
            document.add_section(WD_SECTION.NEW_PAGE)
    report: dict = {}
    _strip_running_furniture(document, _page_break_paragraphs(document), report)
    assert report["items"], "the footers should have been recognised"
    assert {item["side"] for item in report["items"]} == {"footer"}
