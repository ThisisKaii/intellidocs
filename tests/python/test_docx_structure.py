"""DOCX structure fidelity: content controls, text boxes, merged cells, page breaks, tabs."""
from pathlib import Path
import re
import sys

import docx
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches

REPO_ROOT = Path(__file__).resolve().parents[2]
ML_SRC = REPO_ROOT / "ml" / "src"
if str(ML_SRC) not in sys.path:
    sys.path.insert(0, str(ML_SRC))

from converter import _docx_document_to_result  # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _html(doc: docx.Document) -> str:
    return _docx_document_to_result(doc)["html"]


def _text(html: str) -> str:
    return re.sub(r"<[^>]+>", "", html)


def _append_before_sectpr(doc: docx.Document, element: object) -> None:
    sect_pr = doc.element.body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(element)
    else:
        doc.element.body.append(element)


class TestContentControls:
    """A block w:sdt (table of contents, cover page) is walked like the body."""

    def test_block_content_control_paragraphs_are_emitted(self) -> None:
        doc = docx.Document()
        sdt = parse_xml(
            f'<w:sdt xmlns:w="{W_NS}"><w:sdtContent>'
            '<w:p><w:r><w:t>TITLE PAGE</w:t></w:r></w:p>'
            '<w:p><w:r><w:t>APPROVAL SHEET</w:t></w:r></w:p>'
            "</w:sdtContent></w:sdt>"
        )
        _append_before_sectpr(doc, sdt)
        text = _text(_html(doc))
        assert "TITLE PAGE" in text and "APPROVAL SHEET" in text

    def test_nested_and_inline_content_controls(self) -> None:
        doc = docx.Document()
        paragraph = doc.add_paragraph("Before ")
        inline = parse_xml(f'<w:sdt xmlns:w="{W_NS}"><w:sdtContent><w:r><w:t>inline control</w:t></w:r></w:sdtContent></w:sdt>')
        paragraph._p.append(inline)
        tracked = parse_xml(f'<w:ins xmlns:w="{W_NS}" w:id="1" w:author="a"><w:r><w:t> tracked</w:t></w:r></w:ins>')
        paragraph._p.append(tracked)
        assert "inline control" in _html(doc) and "tracked" in _html(doc)


class TestTextBoxes:
    """Text boxes are emitted once even though AlternateContent stores them twice."""

    def test_text_box_text_appears_once(self) -> None:
        doc = docx.Document()
        paragraph = doc.add_paragraph("Anchor line")
        alt = parse_xml(
            f'<mc:AlternateContent xmlns:mc="{MC_NS}" xmlns:w="{W_NS}">'
            '<mc:Choice Requires="wps"><w:txbxContent><w:p><w:r><w:t>Boxed words</w:t></w:r></w:p></w:txbxContent></mc:Choice>'
            '<mc:Fallback><w:txbxContent><w:p><w:r><w:t>Boxed words</w:t></w:r></w:p></w:txbxContent></mc:Fallback>'
            "</mc:AlternateContent>"
        )
        run = paragraph.add_run()
        run._r.append(alt)
        html = _html(doc)
        assert html.count("Boxed words") == 1
        assert html.index("Anchor line") < html.index("Boxed words")


class TestMergedCells:
    """gridSpan becomes colspan; vertical merges are flattened (no rowspan) and never duplicate text."""

    def _table_doc(self) -> docx.Document:
        doc = docx.Document()
        table = doc.add_table(rows=3, cols=3)
        for r in range(3):
            for c in range(3):
                table.cell(r, c).text = f"r{r}c{c}"
        table.cell(0, 0).merge(table.cell(0, 1))  # horizontal
        table.cell(1, 2).merge(table.cell(2, 2))  # vertical
        return doc

    def test_colspan_is_emitted_and_rowspan_is_flattened(self) -> None:
        html = _html(self._table_doc())
        assert 'colspan="2"' in html and "rowspan" not in html

    def test_merged_text_is_not_duplicated(self) -> None:
        text = _text(_html(self._table_doc()))
        assert text.count("r0c0") == 1 and text.count("r1c2") == 1

    def test_vertical_merge_continuation_is_an_empty_cell_without_the_shared_border(self) -> None:
        html = _html(self._table_doc())
        assert "border-top:none" in html and "border-bottom:none" in html
        assert html.count("<td") + html.count("<th") == 8  # every grid cell of the 3 rows exists (9 - 1 horizontal merge)

    def test_table_carries_column_sizes_when_columns_match_cells(self) -> None:
        doc = docx.Document()
        doc.add_table(rows=2, cols=3)
        html = _html(doc)
        match = re.search(r'data-column-size="([\d.,]+)"', html)
        assert match and len(match.group(1).split(",")) == 3
        assert abs(sum(float(v) for v in match.group(1).split(",")) - 100) < 0.1

    def test_column_sizes_omitted_when_merges_hide_the_column_count(self) -> None:
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1))
        assert "data-column-size" not in _html(doc)


class TestPageBreaks:
    """Page-break-before headings and next-page sections start new pages once."""

    def test_page_break_before_emits_a_break(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Intro")
        chapter = doc.add_paragraph("CHAPTER I")
        chapter.paragraph_format.page_break_before = True
        html = _html(doc)
        assert html.count("data-page-break") == 1
        assert html.index("data-page-break") < html.index("CHAPTER I")

    def test_explicit_break_plus_page_break_before_is_one_break(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Intro")
        doc.add_paragraph().add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
        chapter = doc.add_paragraph("CHAPTER I")
        chapter.paragraph_format.page_break_before = True
        assert _html(doc).count("data-page-break") == 1

    def test_next_page_section_breaks_become_page_breaks(self) -> None:
        doc = docx.Document()
        for index in range(3):
            doc.add_paragraph(f"page {index}")
            doc.add_section(WD_SECTION.NEW_PAGE)
        # Three pages have two boundaries; the trailing section break has no content after it.
        assert _html(doc).count("data-page-break") == 2


class TestTabs:
    """Tabs are inline nodes carrying the paragraph's stops; space runs keep their width."""

    def test_tab_becomes_a_node_with_the_paragraphs_stops(self) -> None:
        doc = docx.Document()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run("TITLE PAGE\ti")
        html = _html(doc)
        assert "<span data-tab></span>" in html
        assert "right" in html and "576.0" in html and 'data-default-tab="48.0"' in html

    def test_paragraph_without_tabs_has_no_tab_attributes(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("plain")
        assert "data-tab" not in _html(doc)

    def test_runs_of_spaces_survive_as_nbsp(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Member" + " " * 8 + "Member")
        assert " " * 7 in _html(doc)


class TestParagraphMarkSize:
    """The paragraph mark sizes blank lines only, never a paragraph that has text."""

    def _mark_size(self, paragraph: docx.text.paragraph.Paragraph, half_points: int) -> None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(half_points))
        r_pr = OxmlElement("w:rPr")
        r_pr.append(sz)
        paragraph._p.get_or_add_pPr().append(r_pr)

    def test_text_paragraph_ignores_a_larger_paragraph_mark(self) -> None:
        doc = docx.Document()
        doc.styles["Normal"].font.size = docx.shared.Pt(11)
        paragraph = doc.add_paragraph("Body text without its own size")
        self._mark_size(paragraph, 26)
        assert re.search(r'<p style="[^"]*font-size:11pt', _html(doc))

    def test_blank_paragraph_uses_its_mark_size(self) -> None:
        doc = docx.Document()
        doc.styles["Normal"].font.size = docx.shared.Pt(11)
        blank = doc.add_paragraph()
        self._mark_size(blank, 26)
        assert re.search(r'empty-spacer" style="[^"]*font-size:13pt', _html(doc))
